import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from ec_customer_workbooks import export_customers, split_customers


def example():
    scenarios = []
    for line_key, prca, request_id in [(1, '100', 'REQ-1'), (2, '100', 'REQ-2'), (3, '200', 'REQ-3')]:
        for tier in range(1, 11):
            scenarios.append({'line_key': line_key, 'prca_num': prca, 'request_id': request_id,
                              'candidate_tier': tier, 'item': str(line_key), 'item_dsc': '=DO_NOT_EXECUTE()',
                              'current_tier': 3,
                              'requested_price': '10.2500', 'requested_uom': 'EA', 'sell_uom': 'EA',
                              'tier_cost': '7.00', 'requested_margin_at_candidate_cost': '0.317073',
                              'validation_errors': [], 'impact_exceptions': ['NO_IPC_ACTIVITY'],
                              'grouped_lvl2': 'GAUZE'})
    groups = [{'prca': '100', 'request_id': 'REQ-1', 'category': 'GAUZE',
        'current_tier': 3, 'analyst_benchmark_tier': 3,
        'candidate_tier': tier, 'tier_target_margin': '0.12',
        'baseline_margin': '0.20', 'tier_only_margin': '0.20',
        'requested_price_margin': '0.317073', 'distance_from_tier_target': '0.197073',
        'suggested_for_review': tier == 7, 'requested_line_count': 1,
        'exception_line_count': 0, 'impact_unavailable_line_count': 1,
        'annualized_gp_change': '0', 'annualized_internal_cost_change': '0',
        'annualized_cost_change': '0', 'annualized_customer_spend_change': '0',
        'screening_flags': [], 'exceptions': []} for tier in range(1, 11)]
    activity = [{'prca_num': '100', 'item_e1_num': str(item), 'item_dsc': f'Item {item}',
        'grouped_lvl2': 'GAUZE', 'ipc_sell_uom': 'EA', 'reference_sell_uom': 'EA',
        'qty_3_mth': 2, 'sls_3_mth': 20, 'current_tier': 3,
        'prvt_brnd_flg': 'Y', 'attachment_matches': 1,
        **{f'tier_{tier}': tier + 1 for tier in range(1, 11)}} for item in (1, 2)]
    request_tiers = [{'request_id': 'REQ-1', 'prca': '100', 'category': 'GAUZE',
        'candidate_tier': tier, 'current_tier': 3, 'analyst_benchmark_tier': 3,
        'request_item_count': 2, 'requested_unit_sales_proxy': '20.50',
        'requested_unit_cost_proxy': '14.00', 'requested_item_margin': '0.317073',
        'tier_target_margin': '0.12', 'distance_from_tier_target': '0.197073',
        'suggested_for_review': tier == 7, 'screening_flags': [], 'exceptions': []}
        for tier in range(1, 11)]
    return {'mode': 'INTERNAL_REVIEW_ONLY_NOT_APPROVED_NOT_LOADABLE', 'lines': [{}, {}, {}],
            'scenarios': scenarios, 'groups': groups, 'activity': activity,
            'request_tier_scenarios': request_tiers, 'source_hash': 'test-source', 'reference_hash': 'test-ref',
            'excluded_unpriced_lines': [{'request_id': 'REQ-1', 'account_provided': '20046631',
                'prca_num': '100', 'item': '99',
                'source_item_description': 'Unpriced extra item', 'requested_uom': 'EA',
                'requested_price': '', 'price_exclusion_reason': 'MISSING_REQUESTED_PRICE',
                'source_errors': ['INVALID_REQUESTED_PRICE']}],
            'analyst_tier_assumption': 'Current attached EC tier is assumed to be analyst assignment'}


class CustomerWorkbookTests(unittest.TestCase):
    def test_one_file_per_customer_combines_requests_without_cross_customer_leakage(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary) / 'customers'
            manifest = export_customers(example(), folder)
            self.assertEqual(manifest['customers'], 2)
            self.assertEqual(manifest['requested_lines'], 3)
            self.assertEqual(manifest['tier_scenarios'], 30)
            with (folder / 'Customer_100_Review.xlsx').open('rb') as source:
                workbook = load_workbook(source)
                try:
                    items = workbook['Requested Items']
                    self.assertEqual(items.max_row, 5)
                    self.assertEqual({items.cell(row, 1).value for row in (4, 5)}, {'REQ-1', 'REQ-2'})
                    excluded = workbook['Excluded Unpriced Lines']
                    self.assertEqual(excluded.max_row, 4)
                    self.assertEqual(excluded.cell(4, 1).value, 'REQ-1')
                    self.assertEqual(excluded.cell(4, 2).value, '20046631')
                    self.assertEqual(excluded.cell(4, 3).value, '100')
                    self.assertEqual(excluded.cell(4, 4).value, '99')
                    self.assertEqual(excluded.cell(4, 8).value, 'MISSING_REQUESTED_PRICE')
                    summary_rows = list(workbook['Customer Summary'].iter_rows(
                        min_row=4, values_only=True))
                    self.assertIn(('Unpriced lines excluded from tier analysis', 1), summary_rows)
                    self.assertEqual(items.cell(4, 7).value, '10.2500')
                    self.assertEqual(items.cell(4, 3).data_type, 's')
                    self.assertEqual(workbook['Tier Costs'].cell(4, 4).value, 7)
                    self.assertEqual(workbook['Tier Margins'].cell(4, 4).number_format, '0.0%;[Red](0.0%)')
                    self.assertIn('tier-only margin', workbook['Requested Category Review']['A1'].value)
                    self.assertIn('requested-price scenario margin', workbook['Requested Category Review']['A1'].value)
                    self.assertIn('requested-price impacts are a separate scenario',
                                  workbook['Requested Category Review']['A1'].value)
                    benchmarks = workbook['Requested Tier Benchmarks']
                    self.assertEqual(benchmarks.cell(10, 9).value, '0.12')
                    self.assertAlmostEqual(benchmarks.cell(10, 10).value, 7 / 0.88, places=2)
                    rules = [rule for group in benchmarks.conditional_formatting._cf_rules.values()
                             for rule in group]
                    formulas = [formula for rule in rules for formula in rule.formula]
                    self.assertTrue(any('$K4=FALSE' in formula for formula in formulas))
                    self.assertTrue(any('$L4=TRUE' in formula for formula in formulas))
                    self.assertTrue(any('$G4=$O4' in formula for formula in formulas))
                    for sheet_name in ('Customer Summary', 'Requested Items', 'Tier Costs',
                        'Tier Margins', 'Tier Impact', 'Request Tier Review',
                        'Requested Category Review', 'Existing Category Items', 'Category Item Tier Costs',
                        'Requested Tier Benchmarks', 'Exceptions', 'Analyst Notes', 'Source Notes'):
                        with self.subTest(sheet=sheet_name):
                            self.assertGreater(len(workbook[sheet_name].conditional_formatting), 0)
                    request_tier = workbook['Request Tier Review']
                    self.assertEqual(request_tier.cell(4, 4).value, 3)
                    category = workbook['Requested Category Review']
                    self.assertEqual(category.cell(4, 5).value, 3)
                    benchmarks = workbook['Requested Tier Benchmarks']
                    self.assertEqual(benchmarks.cell(4, 14).value, 3)
                    self.assertEqual(workbook['Existing Category Items'].max_row, 5)
                    self.assertEqual(workbook['Category Item Tier Costs'].cell(4, 7).value, 2)
                finally:
                    workbook.close()
            with (folder / 'Customer_Index.xlsx').open('rb') as source:
                index = load_workbook(source)
                try:
                    self.assertEqual(index['Customers'].cell(4, 7).hyperlink.target,
                                     'Customer_100_Review.xlsx')
                finally:
                    index.close()
            with self.assertRaises(ValueError):
                export_customers(example(), folder)

    def test_customer_with_only_unpriced_lines_still_gets_a_workbook(self):
        report = example()
        report['lines'] = []
        report['scenarios'] = []
        report['excluded_unpriced_lines'] = [{
            'request_id': 'REQ-ONLY-UNPRICED', 'prca_num': '300', 'item': '77',
            'requested_price': '0', 'price_exclusion_reason': 'ZERO_REQUESTED_PRICE',
            'source_errors': ['INVALID_REQUESTED_PRICE'],
        }]
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary) / 'customers'
            manifest = export_customers(report, folder)
            self.assertEqual(manifest['customers'], 1)
            self.assertEqual(manifest['requested_lines'], 0)
            self.assertEqual(manifest['excluded_unpriced_lines'], 1)
            with (folder / 'Customer_300_Review.xlsx').open('rb') as source:
                workbook = load_workbook(source)
                try:
                    self.assertEqual(workbook['Requested Items'].max_row, 3)
                    self.assertEqual(workbook['Excluded Unpriced Lines'].cell(4, 1).value,
                                     'REQ-ONLY-UNPRICED')
                finally:
                    workbook.close()

    def test_unresolved_customers_and_incomplete_tier_sets_fail_closed(self):
        report = example()
        report['scenarios'][0]['prca_num'] = None
        with self.assertRaises(ValueError):
            split_customers(report)
        report = example()
        report['scenarios'].pop()
        with self.assertRaises(ValueError):
            split_customers(report)


if __name__ == '__main__':
    unittest.main()