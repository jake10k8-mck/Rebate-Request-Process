import io
import unittest

from openpyxl import Workbook

from ec_customer_workbooks import add_item_level_conditional_formatting
from ec_manual_case_review import CASES, generate_html, make_manifest, normalized_title, parse_diagnostics_attachment, resolve_smartsheet_account


class ManualCaseReviewTests(unittest.TestCase):
    def test_expected_request_scope_and_separate_analyst_benchmark(self):
        self.assertEqual([case['request_id'] for case in CASES], ['MBR-0011485', 'MBR-0011502'])
        self.assertEqual({case['analyst_tier'] for case in CASES}, {7})
        self.assertEqual(CASES[0]['resolved_prca'], '90586688')

    def test_attachment_selects_request_sheet_and_uses_future_price(self):
        workbook = Workbook()
        detail = workbook.active
        detail.title = 'Item Detail'
        detail.append(['ITEM_E1_NUM', 'ITEM_DSC', 'CTLG_NUM', 'SELL_UOM'])
        detail.append([123456, 'Internal lookup row', 'x', 'EA'])
        request = workbook.create_sheet('Rebate Form')
        request.append(['E1 Item Number', 'Item Description', 'Sell UOM', 'Future Sell Price',
                        'Current Sell Price', 'Current Cost'])
        request.append([654321, 'Requested item', 'BX', '2.12', '3.29', '2.63'])
        content = io.BytesIO()
        workbook.save(content)
        workbook.close()
        lines = parse_diagnostics_attachment(content.getvalue())
        self.assertEqual(len(lines), 1)
        self.assertEqual(str(lines[0]['ITEM_E1_NUM']), '654321')
        self.assertEqual(str(lines[0]['REQ_PRICE']), '2.12')
        self.assertEqual(lines[0]['REQ_UOM'], 'BX')
        self.assertEqual(lines[0]['SOURCE_WORKSHEET'], 'Rebate Form')
    def test_normalizes_header_labels_without_disclosing_values(self):
        self.assertEqual(normalized_title('Future_Sell Price'), 'FUTURE SELL PRICE')

    def test_item_tier_sheet_formats_each_row_from_its_own_values(self):
        workbook = Workbook()
        sheet = workbook.active
        headers = ['Item', 'Requested Price', 'Tier', 'Tier Cost', 'Tier Target Margin',
                   'Benchmark Price', 'Price > Cost', 'Price Meets Benchmark', 'Current Tier',
                   'Analyst Tier 7', 'Request-price Suggested Tier', 'Category Suggested Tier',
                   'Tier 10 Review Flag', 'Exceptions']
        sheet.append(['Item tier scenarios'])
        sheet.append([])
        sheet.append(headers)
        sheet.append(['100', 4.01, 7, 3.45, 0.12, 3.92, True, True, 7, None, 7, None, False, None])
        add_item_level_conditional_formatting(sheet)
        formulas = [formula for group in sheet.conditional_formatting._cf_rules.values()
                    for rule in group for formula in rule.formula]
        self.assertTrue(any('$G4=FALSE' in formula for formula in formulas))
        self.assertTrue(any('$H4=TRUE' in formula for formula in formulas))
        self.assertTrue(any('$C4=$K4' in formula for formula in formulas))

    def test_accepts_both_prca_aliases_and_rejects_conflicts(self):
        self.assertEqual(resolve_smartsheet_account({'PRCA': 90586438.0}, '90586438'), '90586438')
        self.assertEqual(resolve_smartsheet_account(
            {'Customer Account Number/PRCA': '90586438'}, '90586438'), '90586438')
        with self.assertRaises(ValueError):
            resolve_smartsheet_account({'PRCA': '1', 'Customer Account Number/PRCA': '2'}, '90586438')

    def test_summary_scopes_each_benchmark_to_named_category(self):
        report = {'scenarios': []}
        category_rows = []
        for line_number, (case, category) in enumerate([
            (CASES[0], 'DIAGNOSTIC'), (CASES[0], 'GLOVES - NITRILE'),
            (CASES[1], 'GLOVES - NITRILE')], start=1):
            report['scenarios'].append({'line_key': line_number, 'request_id': case['request_id'],
                'candidate_tier': 1, 'grouped_lvl2': category, 'item': str(line_number),
                'item_dsc': None, 'requested_uom': 'EA', 'requested_price': '10',
                'sell_uom': 'EA', 'current_tier': 8, 'current_tier_cost': '7',
                'validation_errors': [], 'impact_exceptions': []})
            for tier in range(1, 11):
                category_rows.append({'request_id': case['request_id'], 'prca': case['resolved_prca'],
                    'category': category, 'candidate_tier': tier, 'current_tier': 8,
                    'exceptions': ['REQUEST_PRICE_DOES_NOT_EXCEED_TIER_COST'] if tier == 1 else [],
                    'screening_flags': [], 'suggested_for_review': tier == 7,
                    'baseline_sales_3m': '100', 'baseline_gp_3m': '20', 'baseline_margin': '0.2',
                    'tier_target_margin': '0.12', 'tier_only_gp_3m': '20', 'tier_only_margin': '0.2',
                    'requested_price_sales_3m': '100', 'requested_price_gp_3m': '20',
                    'requested_price_margin': '0.2', 'distance_from_tier_target': '0.08',
                    'annualized_gp_change': '0', 'annualized_internal_cost_change': '0',
                    'annualized_customer_spend_change': '0', 'assumed_sell_uom_lines': []})
        manifest = make_manifest({**report, 'evaluated_at': 'now', 'source_hash': 's',
            'reference_hash': 'r', 'assumptions': []}, category_rows, [])
        diagnostic = next(case for case in manifest['cases'] if case['request_id'] == CASES[0]['request_id'])
        self.assertEqual(diagnostic['category'], 'DIAGNOSTIC')
        self.assertEqual(diagnostic['suggested_tier'], 7)
        self.assertEqual(diagnostic['other_request_categories_excluded'], ['GLOVES - NITRILE'])
        self.assertIn('Tier 1: REQUEST_PRICE_DOES_NOT_EXCEED_TIER_COST', diagnostic['exceptions'])
        self.assertIn('Account 90586438, PRCA 90586688', generate_html(manifest))

if __name__ == '__main__':
    unittest.main()