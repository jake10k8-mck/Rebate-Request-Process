"""Build manual two-case analyst review workbooks; no approval or email actions."""

import argparse
from collections import Counter
import hashlib
import html
import json
import os
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import Workbook
from ec_category_review import (TIER_TARGET_MARGINS, category_scenarios,
                                requested_item_tier_scenarios, tier_benchmark_price)
from ec_customer_workbooks import (add_activity_data_bars, add_context_status_conditional_formatting,
    add_item_level_conditional_formatting, add_review_row_conditional_formatting, add_sheet,
    add_tier_matrix_conditional_formatting)
from ec_etl_preview import build_preview, connect_from_environment, json_value
from ec_review import calculate_scenarios, group_scenarios, number, numeric_id, prepare_requests, read_references
from ec_test_inputs import fetch
from ec_workflow import utc_now


CASES = (
    {'request_id': 'MBR-0011485', 'account_number': '90586438',
     'resolved_prca': '90586688', 'analyst_tier': 7, 'category_contains': 'DIAGNOSTIC'},
    {'request_id': 'MBR-0011502', 'account_number': '62040942',
     'resolved_prca': '62040942', 'analyst_tier': 7, 'category_contains': 'NITRILE'},
)
SMARTSHEET_STAGES = (
    ('landing', '727058820845444'),
    ('intermediate', '6992940959223684'),
    ('completed', '7801368878206852'),
    ('rejected', '4985786021662596'),
)


def normalized_title(value):
    return ' '.join(re.sub(r'[_\s]+', ' ', str(value or '').strip().upper()).split())


def resolve_smartsheet_account(cells: dict, expected: str) -> str:
    values = [numeric_id(cells.get(key)) for key in ('Customer Account Number/PRCA', 'PRCA')
              if cells.get(key) is not None and str(cells.get(key)).strip()]
    if not values or None in values or len(set(values)) != 1:
        raise ValueError('Smartsheet PRCA fields are missing, invalid, or conflicting')
    account = values[0]
    if account != expected:
        raise ValueError('Smartsheet account does not match the verified benchmark')
    return account


def parse_diagnostics_attachment(content: bytes, request_id: str | None = None) -> list[dict]:
    import io
    from openpyxl import load_workbook

    workbook = load_workbook(content if hasattr(content, 'read') else io.BytesIO(content),
                              read_only=True, data_only=True)
    try:
        found = []
        incomplete_headers = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            for index, row in enumerate(rows[:100]):
                titles = [normalized_title(value) for value in row]
                if 'E1 ITEM NUMBER' not in titles and 'ITEM E1 NUM' not in titles:
                    continue
                columns = {}
                aliases = {'item': {'E1 ITEM NUMBER', 'ITEM E1 NUM'},
                           'uom': {'SELL UOM', 'REQ UOM', 'REQUESTED UOM'},
                           'price': {'FUTURE SELL PRICE', 'REQUESTED SELL PRICE', 'REQ PRICE'}}
                for field, options in aliases.items():
                    matches = [position for position, name in enumerate(titles) if name in options]
                    if len(matches) != 1:
                        incomplete_headers.append(f'{sheet.title}: missing or ambiguous {field}')
                        break
                    columns[field] = matches[0]
                if len(columns) != len(aliases):
                    continue
                result = []
                for row_number, values in enumerate(rows[index + 1:], start=index + 2):
                    cells = [values[columns[field]] if columns[field] < len(values) else None
                             for field in ('item', 'uom', 'price')]
                    if all(value is None or not str(value).strip() for value in cells):
                        continue
                    result.append({'row_id': request_id or CASES[0]['request_id'], 'ITEM_E1_NUM': cells[0],
                                   'ITEM_DSC': None, 'REQ_UOM': cells[1], 'REQ_PRICE': cells[2],
                                   'SOURCE_WORKSHEET': sheet.title, 'SOURCE_ROW_NUMBER': row_number})
                if result:
                    found.append(result)
                break
        if len(found) != 1:
            if not found and incomplete_headers:
                raise ValueError('No complete request worksheet: ' + '; '.join(incomplete_headers))
            raise ValueError('Diagnostics attachment must contain exactly one populated request worksheet')
        return found[0]
    finally:
        workbook.close()


def get_diagnostics_from_smartsheet(env_file: Path | None, case_spec: dict | None = None) -> tuple[dict, list[dict]]:
    case_spec = case_spec or CASES[0]
    token = os.environ.get('SMARTSHEET_API_TOKEN')
    if not token and env_file:
        from dotenv import dotenv_values
        token = dotenv_values(env_file).get('SMARTSHEET_API_TOKEN')
    if not token:
        raise ValueError('Verified Diagnostics attachment is missing; configure approved local Smartsheet access')
    def get_json(path):
        request = Request('https://api.smartsheet.com/2.0/' + path,
                          headers={'Authorization': 'Bearer ' + token})
        with urlopen(request, timeout=90) as response:
            return json.load(response)
    matches = []
    for stage, sheet_id in SMARTSHEET_STAGES:
        sheet = get_json(f'sheets/{sheet_id}')
        titles = {column['id']: column['title'] for column in sheet['columns']}
        for row in sheet['rows']:
            cells = {titles.get(cell['columnId'], ''):
                     (cell.get('value') if cell.get('value') is not None else cell.get('displayValue'))
                     for cell in row['cells']}
            if case_spec['request_id'] in [str(cells.get(name) or '') for name in ('Row ID', 'Original Row ID')]:
                matches.append((stage, sheet_id, row, cells))
    if len(matches) != 1:
        raise ValueError(f'Diagnostics Smartsheet request matched {len(matches)} stages; expected exactly one')
    stage, sheet_id, row, cells = matches[0]
    account = resolve_smartsheet_account(cells, case_spec['account_number'])
    attachment_list = get_json(f'sheets/{sheet_id}/rows/{row["id"]}/attachments')
    attachments = [entry for entry in attachment_list.get('data', [])
                   if entry.get('attachmentType') == 'FILE' and str(entry.get('name', '')).lower().endswith('.xlsx')]
    if len(attachments) != 1:
        raise ValueError('Diagnostics request must have exactly one XLSX attachment')
    attachment = get_json(f'sheets/{sheet_id}/attachments/{attachments[0]["id"]}')
    if urlparse(attachment['url']).scheme != 'https':
        raise ValueError('Attachment download requires HTTPS')
    with urlopen(attachment['url'], timeout=90) as response:
        content = response.read(20 * 1024 * 1024 + 1)
    if len(content) > 20 * 1024 * 1024:
        raise ValueError('Attachment exceeds the manual-review size limit')
    header = {'ROW_ID': case_spec['request_id'], 'CUST_NUM_PRCA': account,
              'CREATED_DT': cells.get('Created'), 'BUS_PLTFRM_DSC': cells.get('Sales Segment'),
              'SOURCE_STAGE': stage, 'SOURCE_SHEET_ID': sheet_id,
              'SOURCE_SHEET_ROW_ID': str(row['id']), 'SOURCE_ATTACHMENT': attachments[0]['name'],
              'SOURCE_ATTACHMENT_SHA256': hashlib.sha256(content).hexdigest()}
    return header, parse_diagnostics_attachment(content, case_spec['request_id'])


def resolve_case_mapping(connection, case_spec: dict) -> dict:
    matches = fetch(connection,
        'SELECT DISTINCT CUSTOMER.CUST_E1_NUM, CUSTOMER.PRCA_NUM, PEX.BUS_PLTFRM '
        'FROM PRD_PL_MMS_EDW_DB.EDWRPT.V_DIM_CUST_CURR CUSTOMER '
        'LEFT JOIN PRD_SBX_MMS_DB.PRC.PEX_PRCA_VIEW PEX ON PEX.PRCA_NUM = CUSTOMER.PRCA_NUM '
        'WHERE CUSTOMER.SYS_PLTFRM = %s AND CUSTOMER.CUST_TYPE_CD IN (%s, %s, %s) '
        'AND CUSTOMER.CUST_E1_NUM = %s',
        ('E1', 'B', 'S', 'X', int(case_spec['account_number'])))
    identities = {(numeric_id(row.get('cust_e1_num')), numeric_id(row.get('prca_num')), row.get('bus_pltfrm'))
                  for row in matches}
    if len(identities) != 1:
        raise ValueError(f'Account {case_spec["account_number"]} has missing or ambiguous customer mapping')
    e1_account, prca, platform = next(iter(identities))
    if e1_account != case_spec['account_number'] or not prca or platform != 'EC':
        raise ValueError('Request account does not map uniquely to a confirmed EC customer')
    if case_spec.get('resolved_prca') and case_spec['resolved_prca'] != prca:
        raise ValueError('Resolved PRCA differs from the known benchmark')
    return {**case_spec, 'resolved_prca': prca, 'bus_pltfrm': platform}


def read_requested_inputs(connection, env_file: Path | None, case_specs: tuple[dict, ...] | list[dict]) -> dict:
    headers, lines = [], []
    for case_spec in case_specs:
        if case_spec['request_id'] == CASES[0]['request_id']:
            header, requested = get_diagnostics_from_smartsheet(env_file, case_spec)
        else:
            requested = fetch(connection,
                'SELECT REQUEST_ID, ACCOUNT_NUMBER, ITEM_E1_NUM, UOM_PROVIDED, FUTURE_SELL_PRICE, CREATED_DATE '
                'FROM PRD_SBX_MMS_DB.PRC.MCKB_RBT_TEST_OCT_2026 WHERE REQUEST_ID = %s ORDER BY ITEM_E1_NUM',
                (case_spec['request_id'],))
            if not requested:
                raise ValueError(f'Request {case_spec["request_id"]} is absent from the October source snapshot')
            if any(numeric_id(row.get('account_number')) != case_spec['account_number'] for row in requested):
                raise ValueError('Request account conflicts with source snapshot; refusing inference')
            created_dates = {str(row.get('created_date')) for row in requested}
            if len(created_dates) != 1:
                raise ValueError('Request has conflicting creation dates')
            header = {'ROW_ID': case_spec['request_id'], 'CUST_NUM_PRCA': case_spec['account_number'],
                      'CREATED_DT': requested[0].get('created_date'), 'BUS_PLTFRM_DSC': 'Extended Care'}
            raw_line_keys = [(numeric_id(row.get('item_e1_num')),
                              str(row.get('uom_provided') or '').strip().upper(),
                              str(number(row.get('future_sell_price'))) if number(row.get('future_sell_price')) is not None else '')
                             for row in requested]
            duplicates = Counter(raw_line_keys)
            requested = [{'ROW_ID': case_spec['request_id'], 'ITEM_E1_NUM': row.get('item_e1_num'),
                          'ITEM_DSC': None, 'REQ_UOM': row.get('uom_provided'),
                          'REQ_PRICE': row.get('future_sell_price'), 'SOURCE_DUPLICATE_COUNT': duplicates[key]}
                         for row, key in zip(requested, raw_line_keys)]
        headers.append(header)
        lines.extend(requested)
    preview = build_preview(headers, lines)
    preview['mode'] = 'READ_ONLY_ETL_PREVIEW_NOT_APPROVED_NOT_LOADABLE'
    preview['limitations'].append('Smartsheet attachments and October snapshot inputs are read-only; no source row was changed.')
    preview['limitations'].append('Current references are a replay, not historical costs at the analyst decision date.')
    return preview


def make_manifest(report, category_rows, activity, case_specs: tuple[dict, ...] | list[dict] = CASES):
    cases = []
    for spec in case_specs:
        request_rows = [row for row in category_rows if row['request_id'] == spec['request_id']]
        if spec.get('category_contains'):
            matches = [row for row in request_rows
                       if spec['category_contains'] in str(row.get('category') or '').upper()]
        else:
            available_categories = {row.get('category') for row in request_rows if row.get('category')}
            if len(available_categories) != 1:
                raise ValueError(f'Request {spec["request_id"]} maps to multiple/no categories; pass --category')
            matches = [row for row in request_rows if row.get('category') in available_categories]
        if len(matches) != 10:
            raise ValueError(f'Expected ten {spec["category_contains"]} scenarios for {spec["request_id"]}, got {len(matches)}')
        categories = {row['category'] for row in matches}
        prcas = {row['prca'] for row in matches}
        if len(categories) != 1 or prcas != {spec.get('resolved_prca')}:
            raise ValueError('Resolved PRCA/category differs from analyst benchmark')
        suggested = [row for row in matches if row['suggested_for_review']]
        if len(suggested) > 1:
            raise ValueError('Category review unexpectedly contains multiple suggestions')
        request_scenarios = [row for row in requested_item_tier_scenarios(report)
                             if row['request_id'] == spec['request_id']
                             and row['category'] == next(iter(categories))]
        request_suggested = [row for row in request_scenarios if row['suggested_for_review']]
        if len(request_suggested) > 1:
            raise ValueError('Request-price review unexpectedly contains multiple suggestions')
        scenarios = []
        for row in matches:
            scenarios.append({key: row.get(key) for key in (
                'current_tier', 'candidate_tier', 'baseline_sales_3m', 'baseline_gp_3m',
                'baseline_margin', 'tier_target_margin', 'tier_only_gp_3m', 'tier_only_margin',
                'requested_price_sales_3m', 'requested_price_gp_3m', 'requested_price_margin',
                'distance_from_tier_target', 'annualized_gp_change',
                'annualized_internal_cost_change', 'annualized_customer_spend_change',
                'suggested_for_review', 'assumed_sell_uom_lines', 'screening_flags', 'exceptions')})
        all_request_lines = [row for row in report['scenarios']
                     if row['request_id'] == spec['request_id'] and row['candidate_tier'] == 1]
        selected_category = next(iter(categories))
        lines = [row for row in all_request_lines if row.get('grouped_lvl2') == selected_category]
        other_categories = sorted({str(row.get('grouped_lvl2')) for row in all_request_lines
                       if row.get('grouped_lvl2') != selected_category})
        cases.append({**spec, 'account': spec['account_number'],
            'prca': spec['resolved_prca'], 'category': next(iter(categories)),
            'current_tier': matches[0]['current_tier'],
            'suggested_tier': suggested[0]['candidate_tier'] if suggested else None,
            'request_price_suggested_tier': request_suggested[0]['candidate_tier'] if request_suggested else None,
            'analyst_tier': spec.get('analyst_tier'),
            'status': 'CALCULATED_FOR_REVIEW' if suggested else 'REVIEW_REQUIRED',
            'request_price_status': 'CALCULATED_FOR_REVIEW' if request_suggested else 'REVIEW_REQUIRED',
            'rationale': ('Unique eligible category scenario with current category baseline.' if suggested else
                          'No unique eligible tier suggestion. Review category baseline and exceptions.'),
            'request_price_rationale': ('Unique eligible tier based on requested item prices and candidate costs; '
                'one sell unit per distinct requested item is used because quantities are unavailable.'
                if request_suggested else
                'No unique eligible request-price tier; review requested item costs and exceptions.'),
            'exceptions': sorted({f'Tier {row["candidate_tier"]}: {error}'
                                  for row in matches for error in row['exceptions'] + row['screening_flags']}),
            'source_warnings': sorted({warning for row in lines for warning in row.get('source_warnings', [])}),
            'other_request_categories_excluded': other_categories,
            'customer_sales_context': {
                flag: (str(sum(number(row.get('sls_3_mth')) for row in activity
                               if numeric_id(row.get('prca_num')) == spec['resolved_prca']
                               and row.get('prvt_brnd_flg') == flag))
                       if activity and all(number(row.get('sls_3_mth')) is not None for row in activity
                           if numeric_id(row.get('prca_num')) == spec['resolved_prca']
                           and row.get('prvt_brnd_flg') == flag)
                       and any(numeric_id(row.get('prca_num')) == spec['resolved_prca']
                               and row.get('prvt_brnd_flg') == flag for row in activity) else None)
                for flag in ('Y', 'N')},
            'requested_items': [{'item': row['item'], 'description': row.get('item_dsc'),
                'requested_uom': row['requested_uom'], 'sell_uom': row.get('sell_uom'),
                'uom_assumed_from_sell_uom': not bool(str(row.get('requested_uom') or '').strip()),
                'requested_price': row['requested_price'], 'current_tier': row.get('current_tier'),
                'source_duplicate_count': row.get('source_duplicate_count', 1),
                'current_tier_cost': row.get('current_tier_cost'),
                'analyst_benchmark_cost': next((option['tier_cost'] for option in report['scenarios']
                    if option['request_id'] == spec['request_id'] and option['item'] == row['item']
                    and spec.get('analyst_tier') is not None
                    and option['candidate_tier'] == spec['analyst_tier']), None),
                'validation_errors': row['validation_errors'], 'impact_exceptions': row['impact_exceptions']}
                for row in lines], 'request_scenarios': request_scenarios, 'scenarios': scenarios})
    return {'mode': 'MANUAL_TWO_CASE_REVIEW_NOT_APPROVED', 'recipient': 'jake.tenkate@mckesson.com',
            'evaluated_at': report['evaluated_at'], 'source_hash': report['source_hash'],
            'reference_hash': report['reference_hash'], 'cases': cases,
            'customer_context': [row for row in activity if numeric_id(row.get('prca_num'))
                                in {case['resolved_prca'] for case in case_specs}],
            'assumptions': [
                'Analyst Tier 7 decisions are user-reported historical benchmarks, not forced calculated outcomes.',
                'Current customer/category baseline uses all mapped Tapestry MCKB items with IPC activity.',
                'National-brand customer sales are displayed separately; no projected conversion opportunity was supplied.',
                'Active Tapestry tier targets are exact discrete analyst values: ' + ', '.join(
                    f'Tier {tier}={margin:.0%}' for tier, margin in TIER_TARGET_MARGINS.items()),
                'Historical individual-cost defaults (20% generally, 18% nitrile, 10% vinyl) are context only and do not override the Tapestry tier schedule.',
                'Annualized changes multiply three-month activity by four; these are run-rate estimates, not forecasts.',
                'Missing and UOM-incompatible activity remains unavailable or an exception; it is never imputed as zero.',
                'Current references are as of evaluation time, not a reconstruction of historical analyst costs.',
                'No email, approval write, fixed-price load, or source-system update was performed.'
            ]}


def generate_html(manifest):
    rows = []
    for case in manifest['cases']:
        body = []
        request_body = []
        for scenario in case['request_scenarios']:
            request_body.append('<tr>' + ''.join('<td>' + html.escape(str(
                scenario.get(field) if scenario.get(field) is not None else 'Unavailable')) + '</td>'
                for field in ('candidate_tier', 'request_item_count', 'requested_unit_sales_proxy',
                    'requested_unit_cost_proxy', 'requested_item_margin', 'tier_target_margin',
                    'distance_from_tier_target', 'suggested_for_review', 'screening_flags',
                    'exceptions')) + '</tr>')
        for scenario in case['scenarios']:
            body.append('<tr>' + ''.join('<td>' + html.escape(str(scenario.get(field) if scenario.get(field) is not None else 'Unavailable')) + '</td>' for field in
                ('candidate_tier', 'tier_target_margin', 'baseline_margin', 'tier_only_margin',
                 'distance_from_tier_target', 'requested_price_margin',
                 'annualized_gp_change', 'annualized_internal_cost_change',
                 'annualized_customer_spend_change', 'suggested_for_review', 'exceptions')) + '</tr>')
        items = ''.join('<tr><td>' + '</td><td>'.join(html.escape(str(item.get(field) if item.get(field) is not None else 'Unavailable'))
            for field in ('item', 'description', 'requested_uom', 'sell_uom', 'requested_price', 'current_tier', 'current_tier_cost', 'analyst_benchmark_cost', 'source_duplicate_count')) + '</td></tr>'
            for item in case['requested_items'])
        customer_context = case['customer_sales_context']
        context_text = ('MCKB sales 3m: ' + str(customer_context['Y'] or 'Unavailable')
                        + '; National-brand sales 3m: '
                        + str(customer_context['N'] or 'Unavailable'))
        analyst_label = (f'Tier {case["analyst_tier"]} (historical)' if case.get('analyst_tier') is not None
                         else 'No historical analyst tier supplied')
        duplicate_warning = (' Source snapshot contains identical duplicate rows; they are preserved and flagged.'
                             if case.get('source_warnings') else '')
        rows.append(f'<h2>{html.escape(case["request_id"])} - {html.escape(case["category"])}</h2>'
            f'<p>Account {case["account"]}, PRCA {case["prca"]}. Current tier: {case["current_tier"]}. '
            f'Request-price tier suggestion: {case["request_price_suggested_tier"] or "Unavailable"}; '
            f'customer-category suggestion: {case["suggested_tier"] or "Unavailable"}. '
            'analyst benchmark: ' + html.escape(analyst_label) + '. ' + html.escape(context_text)
            + html.escape(duplicate_warning) + '</p>'
            '<h3>Requested items</h3><table><tr><th>Item</th><th>Description</th><th>Requested UOM</th><th>Sell UOM</th><th>Requested Price</th><th>Current Tier</th><th>Current Cost</th><th>Analyst Benchmark Cost</th><th>Duplicate Source Count</th></tr>' + items + '</table>'
            '<h3>Request-price tier scenarios</h3><table><tr><th>Tier</th><th>Distinct Items</th><th>Sales Proxy</th><th>Cost Proxy</th><th>Requested Item Margin</th><th>Tier Target Margin</th><th>Distance From Target</th><th>Suggested</th><th>Screening Flags</th><th>Exceptions</th></tr>' + ''.join(request_body) + '</table>'
            '<h3>Customer category scenarios</h3><table><tr><th>Tier</th><th>Tier Target Margin</th><th>Baseline Margin</th><th>Tier-only Margin</th><th>Distance From Tier Target</th><th>Requested-price Margin</th><th>Annualized GP Δ</th><th>Annualized Internal Cost Δ</th><th>Annualized Customer Spend Δ</th><th>Suggested for Review</th><th>Exceptions</th></tr>' + ''.join(body) + '</table>'
            f'<p><b>Rationale:</b> {html.escape(case["rationale"])}<br><b>Exceptions:</b> {html.escape("; ".join(case["exceptions"]) or "None")}</p>')
    tier_schedule = ', '.join(f'T{tier}: {margin:.0%}' for tier, margin in TIER_TARGET_MARGINS.items())
    return '<html><head><meta charset="utf-8"><style>body{font-family:Calibri,Arial;color:#222}table{border-collapse:collapse;font-size:11px}th,td{border:1px solid #aaa;padding:5px}th{background:#edf3f5}</style></head><body><h1>Manual EC Tier Review</h1><p>Analyst comparison examples or one-off request review. Suggested tiers are decision support only. No email or load was performed.</p>' + ''.join(rows) + '<p>Active Tapestry tier targets: ' + html.escape(tier_schedule) + '. Category ranking compares requested-price scenario margin to each tier-specific target: requested items with IPC activity use requested price at observed quantity, while other category items retain IPC sales. No volume is assumed for requested items without IPC sales. Tier-only margin remains a separate comparison. Historical individual-cost defaults (20% generally; 18% Nitrile; 10% Vinyl) are context only and do not override this schedule. Existing category business includes mapped Tapestry MCKB items. MCKB and national-brand customer sales are separate context totals, not category margin. Opportunity projection is unavailable and excluded. Annualization is three-month activity × 4, not a forecast. Current reference state may differ from the analyst decision date.</p></body></html>'


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--smartsheet-env-file', type=Path)
    parser.add_argument('--request-id', help='Optional one-off October snapshot request ID')
    parser.add_argument('--account-number', help='Required with --request-id; E1 customer number')
    parser.add_argument('--category', help='Optional category substring if the request maps to multiple categories')
    args = parser.parse_args()
    if bool(args.request_id) != bool(args.account_number):
        parser.error('--request-id and --account-number must be provided together')
    if args.category and not args.request_id:
        parser.error('--category is only valid with --request-id')
    if args.output_dir.exists():
        parser.error('Choose a new output directory; existing review files are never overwritten')
    case_specs = ([{'request_id': args.request_id, 'account_number': numeric_id(args.account_number),
                    'resolved_prca': None, 'analyst_tier': None,
                    'category_contains': args.category.upper() if args.category else None}]
                  if args.request_id else [dict(spec) for spec in CASES])
    if any(spec.get('account_number') is None for spec in case_specs):
        parser.error('--account-number must be a numeric E1 customer number')
    connection = connect_from_environment()
    try:
        case_specs = [resolve_case_mapping(connection, spec) for spec in case_specs]
        preview = read_requested_inputs(connection, args.smartsheet_env_file, case_specs)
        inputs, context = prepare_requests(preview)
        item_refs = read_references(connection, inputs)
        scenarios = calculate_scenarios(context, item_refs)
        base_report = {'mode': 'INTERNAL_REVIEW_ONLY_NOT_APPROVED_NOT_LOADABLE',
            'source_hash': preview['source_hash'], 'evaluated_at': utc_now(), 'lines': context,
            'scenarios': scenarios, 'groups': group_scenarios(scenarios),
            'assumptions': preview['limitations']}
        prcas = sorted({numeric_id(spec['resolved_prca']) for spec in case_specs})
        cursor = connection.cursor()
        try:
            activity_sql = (Path(__file__).parent / 'sql/ec_customer_activity.snowflake.sql').read_text(encoding='utf-8')
            cursor.execute(activity_sql, (json.dumps(prcas),))
            columns = [column[0].lower() for column in cursor.description]
            activity = [dict(zip(columns, row)) for row in cursor.fetchall()]
        finally:
            cursor.close()
        categories = category_scenarios(base_report, activity)
        reference_hash = hashlib.sha256(json.dumps({'item': item_refs, 'activity': activity},
            default=json_value, sort_keys=True).encode('utf-8')).hexdigest()
        report = {**base_report, 'reference_hash': reference_hash}
        manifest = make_manifest(report, categories, activity, case_specs)
    finally:
        connection.close()
    args.output_dir.mkdir(parents=True, exist_ok=False)
    (args.output_dir / 'manual_summary.html').write_text(generate_html(manifest), encoding='utf-8')
    (args.output_dir / 'review_manifest.json').write_text(json.dumps(manifest, default=json_value, indent=2), encoding='utf-8')
    for case in manifest['cases']:
        customer_scenarios = [row for row in scenarios if row['request_id'] == case['request_id']
                              and row.get('grouped_lvl2') == case['category']]
        workbook = Workbook()
        workbook.remove(workbook.active)
        requested_sheet = add_sheet(workbook, 'Requested Items', [
            ('item', 'Item'), ('item_dsc', 'Description'), ('requested_uom', 'Requested UOM'),
            ('sell_uom', 'Sell UOM'), ('requested_price', 'Requested Fixed Price'),
            ('uom_assumed_from_sell_uom', 'UOM Assumed'), ('current_tier', 'Current Tier'),
            ('current_tier_cost', 'Current Tier Cost'), ('validation_errors', 'Validation Exceptions'),
            ('impact_exceptions', 'Impact Exceptions')],
            case['requested_items'], 'Request item price remains unchanged. Blank UOM is assumed sell UOM and explicitly labeled.')
        add_review_row_conditional_formatting(requested_sheet)
        request_sheet = add_sheet(workbook, 'Request Tier Scenarios', [
            ('candidate_tier', 'Candidate Tier'), ('request_item_count', 'Distinct Requested Items'),
            ('current_tier', 'Current Tier'), ('analyst_benchmark_tier', 'Analyst Benchmark Tier'),
            ('requested_unit_sales_proxy', 'Sales Proxy, One Unit per Item'),
            ('requested_unit_cost_proxy', 'Cost Proxy, One Unit per Item'),
            ('requested_item_margin', 'Requested Item Margin'), ('tier_target_margin', 'Tier Target Margin'),
            ('distance_from_tier_target', 'Distance From Target'),
            ('suggested_for_review', 'Suggested for Review'), ('screening_flags', 'Screening Flags'),
            ('exceptions', 'Exceptions')], [dict(row, current_tier=case['current_tier'],
                analyst_benchmark_tier=case['analyst_tier']) for row in case['request_scenarios']],
            'Request-only tier selection uses requested prices and candidate costs; one sell unit per distinct SKU is a proxy because requested quantities are unavailable. Repeated source lines are counted once for this calculation.')
        add_review_row_conditional_formatting(request_sheet)
        category_rows = [dict(row, analyst_benchmark_tier=case['analyst_tier'])
                         for row in case['scenarios']]
        category_sheet = add_sheet(workbook, 'Category Scenarios', [
            ('candidate_tier', 'Candidate Tier'), ('baseline_sales_3m', 'Baseline Sales 3m'),
            ('current_tier', 'Current Tier'), ('analyst_benchmark_tier', 'Analyst Benchmark Tier'),
            ('baseline_gp_3m', 'Baseline GP 3m'), ('baseline_margin', 'Baseline Margin'),
            ('tier_target_margin', 'Tier Target Margin'), ('tier_only_gp_3m', 'Tier-only GP 3m'),
            ('tier_only_margin', 'Tier-only Margin'), ('distance_from_tier_target', 'Distance From Tier Target'),
            ('requested_price_sales_3m', 'Requested Price Sales Scenario 3m'),
            ('requested_price_gp_3m', 'Requested Price GP Scenario 3m'),
            ('requested_price_margin', 'Requested Price Margin'),
            ('annualized_gp_change', 'Annualized GP Change'),
            ('annualized_internal_cost_change', 'Annualized Internal Cost Change'),
            ('annualized_customer_spend_change', 'Annualized Customer Spend Change'),
            ('suggested_for_review', 'Suggested for Review'), ('screening_flags', 'Screening Flags'),
            ('exceptions', 'Exceptions')], category_rows,
            'Green = suggested for review; blue = current tier; gold = historical analyst benchmark tier. Suggestion is not an approval.')
        add_review_row_conditional_formatting(category_sheet)
        benchmark_rows = []
        for item in case['requested_items']:
            item_scenarios = [row for row in customer_scenarios if row['item'] == item['item']]
            tier10_cost = number(next((row.get('tier_cost') for row in item_scenarios
                                       if row['candidate_tier'] == 10), None))
            for scenario in item_scenarios:
                target = TIER_TARGET_MARGINS[scenario['candidate_tier']]
                price, cost, benchmark = (number(scenario.get('requested_price')),
                    number(scenario.get('tier_cost')), number(scenario.get('benchmark_price')))
                calculated_benchmark = tier_benchmark_price(cost, scenario['candidate_tier'])
                if benchmark is not None and calculated_benchmark is not None and benchmark != calculated_benchmark:
                    raise ValueError('SQL and Python benchmark sell prices differ')
                benchmark_rows.append({
                    'request_id': scenario['request_id'], 'item': item['item'],
                    'item_dsc': item.get('description'), 'requested_price': item['requested_price'],
                    'requested_uom': item.get('requested_uom') or item.get('sell_uom'),
                    'uom_assumed_from_sell_uom': item.get('uom_assumed_from_sell_uom'),
                    'candidate_tier': scenario['candidate_tier'], 'tier_cost': scenario.get('tier_cost'),
                    'tier_target_margin': str(target.quantize(Decimal('0.000001'))),
                    'benchmark_price': str(calculated_benchmark) if calculated_benchmark is not None else None,
                    'price_exceeds_cost': price is not None and cost is not None and price > cost,
                    'price_meets_target_benchmark': price is not None and benchmark is not None and price >= benchmark,
                    'current_tier': scenario.get('current_tier'), 'analyst_tier': case['analyst_tier'],
                    'suggested_tier': case['suggested_tier'],
                    'request_price_suggested_tier': case['request_price_suggested_tier'],
                    'tier_10_exception': price is not None and tier10_cost is not None and price <= tier10_cost,
                    'validation_errors': scenario['validation_errors']})
        benchmark_sheet = add_sheet(workbook, 'Requested Tier Benchmarks', [
            ('request_id', 'Request'), ('item', 'Item'), ('item_dsc', 'Description'),
            ('requested_price', 'Requested Price'), ('requested_uom', 'Effective UOM'),
            ('uom_assumed_from_sell_uom', 'UOM Assumed'), ('candidate_tier', 'Tier'),
            ('tier_cost', 'Tier Cost'), ('tier_target_margin', 'Tier Target Margin'),
            ('benchmark_price', 'Sell Price at Target'), ('price_exceeds_cost', 'Price > Cost'),
            ('price_meets_target_benchmark', 'Price Meets Benchmark'), ('current_tier', 'Current Tier'),
            ('analyst_tier', 'Analyst Tier 7'),
            ('request_price_suggested_tier', 'Request-price Suggested Tier'),
            ('suggested_tier', 'Category Suggested Tier'),
            ('tier_10_exception', 'Tier 10 Review Flag'), ('validation_errors', 'Exceptions')],
            benchmark_rows, 'Benchmark price = tier cost ÷ (1 - target margin). Tier 10 threshold is a review flag, not an automatic rejection.')
        add_item_level_conditional_formatting(benchmark_sheet)
        category_activity = [row for row in activity if numeric_id(row.get('prca_num')) == case['prca']
                             and row.get('grouped_lvl2') == case['category']]
        activity_sheet = add_sheet(workbook, 'Existing Category Items', [
            ('item_e1_num', 'Item'), ('item_dsc', 'Description'), ('ipc_sell_uom', 'IPC Sell UOM'),
            ('reference_sell_uom', 'Reference Sell UOM'), ('qty_3_mth', 'Quantity 3m'),
            ('sls_3_mth', 'Sales 3m'), ('current_tier', 'Current Tier'), ('prvt_brnd_flg', 'Private Brand Flag')],
            category_activity, 'Current mapped activity for all category items, not only requested items.')
        add_activity_data_bars(activity_sheet)
        cost_sheet = add_sheet(workbook, 'Category Item Tier Costs', [
            ('item_e1_num', 'Item'), ('item_dsc', 'Description'), ('ipc_sell_uom', 'IPC Sell UOM'),
            ('reference_sell_uom', 'Reference Sell UOM'), ('current_tier', 'Current Tier'),
            *[(f'tier_{tier}', f'Tier {tier} Cost') for tier in range(1, 11)]],
            category_activity, 'Current Tapestry candidate costs for each mapped category item, on sell UOM basis.')
        add_tier_matrix_conditional_formatting(cost_sheet, 'cost')
        customer_activity = [row for row in activity if numeric_id(row.get('prca_num')) == case['prca']]
        summary = []
        for flag, label in [('Y', 'MCKB customer sales 3m'), ('N', 'National brand customer sales 3m')]:
            values = [number(row.get('sls_3_mth')) for row in customer_activity if row.get('prvt_brnd_flg') == flag]
            summary.append({'field': label, 'value': str(sum(values)) if values and all(x is not None for x in values) else None})
        summary.extend([{'field': 'Current tier', 'value': case['current_tier']},
            {'field': 'Request-price suggested tier', 'value': case['request_price_suggested_tier']},
            {'field': 'Customer-category suggested tier', 'value': case['suggested_tier']},
            {'field': 'Analyst historical benchmark', 'value': case['analyst_tier']},
            {'field': 'MCKB customer sales, 3m', 'value': case['customer_sales_context']['Y']},
            {'field': 'National-brand customer sales, 3m', 'value': case['customer_sales_context']['N']},
            {'field': 'Recommendation status', 'value': case['status']},
            {'field': 'Rationale', 'value': case['rationale']},
            {'field': 'Exceptions', 'value': case['exceptions']},
            {'field': 'Opportunity', 'value': 'Not provided; excluded from existing business'}])
        context_sheet = add_sheet(workbook, 'Customer Context', [('field', 'Field'), ('value', 'Value')], summary,
                  'Customer-wide brand sales are context, not category margin denominator.')
        add_context_status_conditional_formatting(context_sheet)
        filename = f'{case["request_id"]}_Manual_Review.xlsx'
        workbook.save(args.output_dir / filename)
        workbook.close()
    print(json.dumps({'output_directory': str(args.output_dir), 'request_ids': [case['request_id'] for case in manifest['cases']],
                      'suggested_tiers': [{'request_id': case['request_id'],
                                           'request_price_suggested_tier': case['request_price_suggested_tier'],
                                           'request_price_status': case['request_price_status'],
                                           'category_suggested_tier': case['suggested_tier'],
                                           'category_status': case['status'],
                                           'current_tier': case['current_tier']}
                                          for case in manifest['cases']], 'email_sent': False}, indent=2))


if __name__ == '__main__':
    main()