"""Generate internal Excel review files, one per resolved customer/PRCA."""

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from ec_category_review import TIER_TARGET_MARGINS, tier_benchmark_price
from ec_review import number, numeric_id
from ec_workflow import utc_now


MONEY = '#,##0.00;[Red](#,##0.00)'
PERCENT = '0.0%;[Red](0.0%)'
NUMBER_FIELDS = {
    'requested_price', 'mac_cost_amt', 'current_tier_cost', 'qty_3_mth', 'sls_3_mth',
    'historical_realized_unit_price', 'tier_cost', 'annualized_cost_change',
    'annualized_customer_spend_change', 'annualized_gross_profit_change_proxy', 'benchmark_price',
}


def excel_value(cell, value, numeric=False, percent=False):
    if isinstance(value, list):
        value = '; '.join(str(entry) for entry in value)
    if value is None:
        return
    if numeric and number(value) is not None:
        cell.value = float(number(value))
        cell.number_format = PERCENT if percent else MONEY
    elif isinstance(value, (int, float)):
        cell.value = value
    else:
        cell.value = str(value)
        cell.data_type = 's'
    cell.alignment = Alignment(vertical='top', wrap_text=True)


def add_sheet(workbook, name, columns, records, note):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.append([note])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    sheet.cell(1, 1).font = Font(name='Calibri', size=11, bold=True, color='36526A')
    sheet.cell(1, 1).alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 42
    for index, (field, label) in enumerate(columns, start=1):
        header = sheet.cell(3, index, label)
        header.font = Font(name='Calibri', bold=True, color='FFFFFF')
        header.fill = PatternFill('solid', fgColor='285B72')
        header.alignment = Alignment(wrap_text=True, vertical='center')
        width = 20
        if any(word in field for word in ('description', 'errors', 'exceptions', 'reason', 'value')):
            width = 42
        sheet.column_dimensions[header.column_letter].width = width
    sheet.row_dimensions[3].height = 32
    for row_number, record in enumerate(records, start=4):
        for column_number, (field, _) in enumerate(columns, start=1):
            numeric = field in NUMBER_FIELDS or field.startswith(('cost_', 'margin_'))
            percent = field.startswith(('margin_', 'target_margin_')) or 'requested_margin' in field
            excel_value(sheet.cell(row_number, column_number), record.get(field), numeric or percent, percent)
    sheet.freeze_panes = 'C4' if len(columns) > 2 else 'A4'
    if records:
        end = sheet.cell(len(records) + 3, len(columns)).coordinate
        table = Table(displayName='T' + name.replace(' ', ''), ref=f'A3:{end}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.print_title_rows = '1:3'
    return sheet


def split_customers(report: dict) -> dict:
    grouped = defaultdict(list)
    owners = defaultdict(set)
    tier_counts = defaultdict(list)
    for row in report['scenarios']:
        prca = numeric_id(row.get('prca_num'))
        if prca is None:
            raise ValueError('Resolve missing customer/PRCA before creating customer files')
        grouped[prca].append(row)
        owners[row['line_key']].add(prca)
        tier_counts[row['line_key']].append(row['candidate_tier'])
    for row in report.get('excluded_unpriced_lines', []):
        prca = numeric_id(row.get('prca_num') or row.get('account_provided'))
        if prca:
            grouped[prca]
    if any(len(values) != 1 for values in owners.values()):
        raise ValueError('One source line resolves to multiple customers')
    if any(sorted(values) != list(range(1, 11)) for values in tier_counts.values()):
        raise ValueError('Every source line must have exactly ten candidate tiers')
    if len(tier_counts) != len(report['lines']):
        raise ValueError('Requested-line counts do not reconcile')
    return dict(sorted(grouped.items()))


def make_customer_workbook(prca: str, scenarios: list[dict], report: dict):
    workbook = Workbook()
    workbook.remove(workbook.active)
    representatives = [row for row in scenarios if row['candidate_tier'] == 1]
    lookup = {(row['line_key'], row['candidate_tier']): row for row in scenarios}
    requests = sorted({row['request_id'] for row in scenarios})
    excluded_unpriced = [row for row in report.get('excluded_unpriced_lines', [])
                         if numeric_id(row.get('prca_num') or row.get('account_provided')) == prca]
    requests = sorted(set(requests) | {row.get('request_id') for row in excluded_unpriced
                                      if row.get('request_id')})
    available = len({row['line_key'] for row in scenarios
                     if row.get('annualized_cost_change') is not None})
    summary = [
        {'field': 'Customer PRCA', 'value': prca},
        {'field': 'Request IDs', 'value': ', '.join(requests)},
        {'field': 'Requested source lines', 'value': len(representatives)},
        {'field': 'Unpriced lines excluded from tier analysis', 'value': len(excluded_unpriced)},
        {'field': 'Tier scenarios', 'value': len(scenarios)},
        {'field': 'Lines with historical impact', 'value': available},
        {'field': 'Reference evaluation timestamp', 'value': report.get('evaluated_at')},
        {'field': 'Tier margin schedule', 'value': ', '.join(f'Tier {tier}: {margin:.0%}' for tier, margin in TIER_TARGET_MARGINS.items())},
        {'field': 'Analyst benchmark assumption', 'value': report.get('analyst_tier_assumption',
            'Current assigned EC tier is assumed to be the analyst-assigned tier; not independently verified')},
        {'field': 'Current coverage', 'value': 'Requested items plus mapped category IPC activity when available'},
        {'field': 'Opportunity', 'value': 'Not supplied; projected conversion sales are not treated as existing sales'},
        {'field': 'Approval state', 'value': 'Not approved; manual notes do not create pipeline decisions or price loads'},
    ]
    summary_sheet = add_sheet(workbook, 'Customer Summary', [('field', 'Field'), ('value', 'Value')], summary,
              'INTERNAL PRICING REVIEW - costs and margins are not rep-facing data.')
    add_context_status_conditional_formatting(summary_sheet)
    item_columns = [
        ('request_id', 'Request'), ('item', 'Item'), ('item_dsc', 'Description'),
        ('requested_uom', 'Requested UOM'), ('sell_uom', 'Sell / Cost UOM'),
        ('requested_price', 'Requested Fixed Price'), ('source_price', 'Exact Source Price'),
        ('mac_cost_amt', 'MAC'), ('current_tier', 'Current Tier'),
        ('current_tier_cost', 'Current Tier Cost'), ('qty_3_mth', 'Quantity 3 Months'),
        ('sls_3_mth', 'Sales 3 Months'), ('validation_errors', 'Validation Exceptions'),
        ('impact_exceptions', 'Impact Exceptions'),
    ]
    items = [dict(row, source_price=row['requested_price']) for row in representatives]
    requested_sheet = add_sheet(workbook, 'Requested Items', item_columns, items,
              'One row per requested source line. Blank impact is unavailable, not zero. Source prices are unchanged.')
    add_review_row_conditional_formatting(requested_sheet)
    add_activity_data_bars(requested_sheet)
    add_sheet(workbook, 'Excluded Unpriced Lines', [
        ('request_id', 'Request'), ('account_provided', 'Source Account'),
        ('prca_num', 'Resolved PRCA'), ('item', 'Item'),
        ('source_item_description', 'Description'), ('requested_uom', 'Requested UOM'),
        ('requested_price', 'Raw Requested Price'),
        ('price_exclusion_reason', 'Excluded From Tier Analysis'),
        ('source_errors', 'Source Validation Flags'),
    ], excluded_unpriced,
        'Blank, zero, negative, or nonnumeric requested prices are excluded from tier analysis but retained here and in the source audit.')
    request_tiers = [row for row in report.get('request_tier_scenarios', [])
                     if numeric_id(row.get('prca')) == prca]
    request_suggestions = {(row.get('request_id'), row.get('category')): row.get('candidate_tier')
                           for row in request_tiers if row.get('suggested_for_review')}
    request_tier_sheet = add_sheet(workbook, 'Request Tier Review', [
        ('request_id', 'Request'), ('category', 'Category'), ('candidate_tier', 'Candidate Tier'),
        ('current_tier', 'Current Assigned Tier'), ('analyst_benchmark_tier', 'Assumed Analyst Tier'),
        ('request_item_count', 'Distinct Requested Items'),
        ('requested_unit_sales_proxy', 'Sales Proxy, One Unit per Item'),
        ('requested_unit_cost_proxy', 'Cost Proxy, One Unit per Item'),
        ('requested_item_margin', 'Requested Item Margin'), ('tier_target_margin', 'Tier Target Margin'),
        ('distance_from_tier_target', 'Distance From Target'),
        ('suggested_for_review', 'Suggested for Review'), ('screening_flags', 'Screening Flags'),
        ('exceptions', 'Exceptions'),
    ], request_tiers,
        'Request-only tier selection uses requested prices and candidate costs. One unit per distinct SKU is a proxy because requested quantities are unavailable; duplicate source lines count once. Current assigned EC tier is the assumed analyst benchmark.')
    add_review_row_conditional_formatting(request_tier_sheet)
    identifiers = [('request_id', 'Request'), ('item', 'Item'), ('item_dsc', 'Description')]
    costs, margins = [], []
    for row in representatives:
        cost_row = {field: row.get(field) for field, _ in identifiers}
        margin_row = dict(cost_row)
        for tier in range(1, 11):
            scenario = lookup[(row['line_key'], tier)]
            cost_row[f'cost_{tier}'] = scenario.get('tier_cost')
            margin_row[f'margin_{tier}'] = scenario.get('requested_margin_at_candidate_cost')
        costs.append(cost_row)
        margins.append(margin_row)
    cost_sheet = add_sheet(workbook, 'Tier Costs', identifiers + [(f'cost_{tier}', f'Tier {tier}') for tier in range(1, 11)],
              costs, 'Tapestry costs represent sell UOM. No tier is selected automatically.')
    add_tier_matrix_conditional_formatting(cost_sheet, 'cost')
    margin_sheet = add_sheet(workbook, 'Tier Margins', identifiers + [(f'margin_{tier}', f'Tier {tier}') for tier in range(1, 11)],
              margins, 'Item margins at requested prices; these are not sales-weighted whole-category margins.')
    add_tier_matrix_conditional_formatting(margin_sheet, 'margin')
    impact_sheet = add_sheet(workbook, 'Tier Impact', [
        ('request_id', 'Request'), ('item', 'Item'), ('grouped_lvl2', 'Category'),
        ('candidate_tier', 'Candidate Tier'), ('annualized_cost_change', 'Annualized Cost Change'),
        ('annualized_customer_spend_change', 'Annualized Customer Spend Change'),
        ('annualized_gross_profit_change_proxy', 'Annualized GP Change Proxy'),
        ('validation_errors', 'Validation Exceptions'), ('impact_exceptions', 'Impact Exceptions'),
    ], scenarios, 'Alternative scenarios, not additive. Annualized means 3-month activity times four, not a forecast.')
    add_review_row_conditional_formatting(impact_sheet)
    category_groups = report.get('category_groups', report.get('groups', []))
    customer_groups = [row for row in category_groups if numeric_id(row.get('prca')) == prca]
    category_records = [{**row, 'analyst_benchmark_tier': row.get('current_tier')}
                        for row in customer_groups]
    customer_activity = [row for row in report.get('activity', [])
                         if numeric_id(row.get('prca_num')) == prca]
    activity_sheet = add_sheet(workbook, 'Existing Category Items', [
        ('item_e1_num', 'Item'), ('item_dsc', 'Description'), ('grouped_lvl2', 'Category'),
        ('ipc_sell_uom', 'IPC Sell UOM'), ('reference_sell_uom', 'Reference Sell UOM'),
        ('qty_3_mth', 'Quantity 3 Months'), ('sls_3_mth', 'Sales 3 Months'),
        ('current_tier', 'Current Tier'), ('prvt_brnd_flg', 'Private Brand Flag'),
        ('attachment_matches', 'Attachment Matches'),
    ], customer_activity, 'Full mapped customer IPC activity returned by the read-only activity query; not restricted to requested items.')
    add_activity_data_bars(activity_sheet)
    tier_cost_rows = [{**row, **{f'cost_{tier}': row.get(f'tier_{tier}') for tier in range(1, 11)}}
                      for row in customer_activity]
    category_cost_sheet = add_sheet(workbook, 'Category Item Tier Costs', [
        ('item_e1_num', 'Item'), ('item_dsc', 'Description'), ('grouped_lvl2', 'Category'),
        ('ipc_sell_uom', 'IPC Sell UOM'), ('reference_sell_uom', 'Reference Sell UOM'),
        ('current_tier', 'Current Tier'),
        *[(f'cost_{tier}', f'Tier {tier}') for tier in range(1, 11)],
    ], tier_cost_rows, 'Current Tapestry candidate costs for all mapped customer IPC items, on sell-UOM basis.')
    add_tier_matrix_conditional_formatting(category_cost_sheet, 'cost')
    category_sheet = add_sheet(workbook, 'Requested Category Review', [
        ('request_id', 'Request'), ('category', 'Category'), ('current_tier', 'Current Tier'),
        ('candidate_tier', 'Candidate Tier'), ('analyst_benchmark_tier', 'Analyst Benchmark Tier'),
        ('tier_target_margin', 'Tier Target Margin'),
        ('baseline_margin', 'Current Category Margin'), ('tier_only_margin', 'Candidate Tier-Only Margin'),
        ('requested_price_margin', 'Requested-Price Scenario Margin'),
        ('distance_from_tier_target', 'Distance From Tier Target'), ('suggested_for_review', 'Suggested for Review'),
        ('requested_line_count', 'Requested Lines'), ('exception_line_count', 'Exception Lines'),
        ('impact_unavailable_line_count', 'Impact Unavailable'),
        ('annualized_gp_change', 'Annualized GP Change'),
        ('annualized_internal_cost_change', 'Annualized Internal Cost Change'),
        ('annualized_cost_change', 'Annualized Cost Change'),
        ('annualized_customer_spend_change', 'Annualized Spend Change'),
        ('screening_flags', 'Screening Flags'), ('exceptions', 'Exceptions'),
    ], category_records, 'Suggestion uses requested-price scenario margin for the full category: requested items with IPC sales use requested price at observed quantity, while non-requested items retain IPC sales. tier-only margin is shown separately. requested-price impacts are a separate scenario from tier-only effects. No volume is assumed for requested items without IPC sales. Current assigned EC tier is used as the assumed analyst benchmark.')
    add_review_row_conditional_formatting(category_sheet)

    requested_benchmarks = []
    for line in representatives:
        for tier in range(1, 11):
            scenario = lookup[(line['line_key'], tier)]
            requested_price = number(line.get('requested_price'))
            tier_cost = number(scenario.get('tier_cost'))
            benchmark_price = number(scenario.get('benchmark_price'))
            requested_benchmarks.append({
                'request_id': line['request_id'], 'item': line.get('item'),
                'item_dsc': line.get('item_dsc'), 'requested_price': line.get('requested_price'),
                'requested_uom': line.get('requested_uom') or scenario.get('sell_uom'),
                'uom_assumed_from_sell_uom': not bool(str(line.get('requested_uom') or '').strip()),
                'candidate_tier': tier, 'tier_cost': scenario.get('tier_cost'),
                'tier_target_margin': str(TIER_TARGET_MARGINS[tier]),
                'benchmark_price': str(tier_benchmark_price(tier_cost, tier)) if tier_benchmark_price(tier_cost, tier) is not None else None,
                'price_exceeds_tier_cost': requested_price is not None and tier_cost is not None and requested_price > tier_cost,
                'price_meets_tier_benchmark': requested_price is not None and benchmark_price is not None and requested_price >= benchmark_price,
                'current_tier': scenario.get('current_tier'),
                'analyst_benchmark_tier': scenario.get('current_tier'),
                'request_price_suggested_tier': request_suggestions.get(
                    (line['request_id'], line.get('grouped_lvl2'))),
                'suggested_tier': next((group.get('candidate_tier') for group in customer_groups
                    if group.get('request_id') == line['request_id'] and group.get('suggested_for_review')), None),
                'validation_errors': scenario['validation_errors'],
                'screening_flags': ['AT_OR_BELOW_TIER_10_COST_REQUIRES_REVIEW']
                    if requested_price is not None and number(lookup[(line['line_key'], 10)].get('tier_cost')) is not None
                    and requested_price <= number(lookup[(line['line_key'], 10)].get('tier_cost')) else [],
            })
    benchmark_sheet = add_sheet(workbook, 'Requested Tier Benchmarks', [
        ('request_id', 'Request'), ('item', 'Requested Item'), ('item_dsc', 'Description'),
        ('requested_price', 'Requested Fixed Price'), ('requested_uom', 'Effective UOM'),
        ('uom_assumed_from_sell_uom', 'UOM Assumed'), ('candidate_tier', 'Candidate Tier'),
        ('tier_cost', 'Tier Cost'), ('tier_target_margin', 'Tier Target Margin'),
        ('benchmark_price', 'Sell Price at Tier Target'), ('price_exceeds_tier_cost', 'Price > Cost'),
        ('price_meets_tier_benchmark', 'Price Meets Benchmark'), ('current_tier', 'Current Tier'),
        ('analyst_benchmark_tier', 'Assumed Analyst Tier'), ('suggested_tier', 'Category Suggested Tier'),
        ('request_price_suggested_tier', 'Request-price Suggested Tier'),
        ('validation_errors', 'Exceptions'), ('screening_flags', 'Screening Flags'),
    ], requested_benchmarks, 'Margin schedule from analyst code: target = -2.22% × tier + 27.22%. Benchmark sell price = tier cost ÷ (1 - target margin).')
    color_benchmark_rows(benchmark_sheet, requested_benchmarks)
    exception_rows = []
    for row in representatives:
        validation = sorted({error for tier in range(1, 11)
                             for error in lookup[(row['line_key'], tier)]['validation_errors']})
        impact = sorted({error for tier in range(1, 11)
                         for error in lookup[(row['line_key'], tier)]['impact_exceptions']})
        if validation or impact:
            exception_rows.append(dict(row, validation_errors=validation, impact_exceptions=impact))
    exception_sheet = add_sheet(workbook, 'Exceptions', identifiers + [('validation_errors', 'Validation Exceptions'),
              ('impact_exceptions', 'Impact Exceptions')], exception_rows,
              'Missing data and unavailable impact are explicit; no source rows are removed.')
    add_review_row_conditional_formatting(exception_sheet)
    categories = sorted({str(row.get('grouped_lvl2') or 'UNMATCHED') for row in representatives})
    analyst_rows = [{'category': category} for category in categories]
    analyst = add_sheet(workbook, 'Analyst Notes', [('category', 'Category'), ('tier', 'Proposed Tier'),
                        ('decision', 'Review Observation'), ('reason', 'Rationale')], analyst_rows,
                        'Manual evaluation notes only. Do not approve category changes from requested-only coverage.')
    add_analyst_note_conditional_formatting(analyst)
    if analyst_rows:
        decisions = DataValidation(type='list', formula1='"Approve,Reject,Needs Review"')
        analyst.add_data_validation(decisions)
        decisions.add(f'C4:C{len(analyst_rows) + 3}')
        tiers = DataValidation(type='whole', operator='between', formula1=1, formula2=10)
        analyst.add_data_validation(tiers)
        tiers.add(f'B4:B{len(analyst_rows) + 3}')
    metadata = [{'field': 'Source hash', 'value': report.get('source_hash')},
                {'field': 'Reference hash', 'value': report.get('reference_hash')}]
    metadata.extend({'field': 'Assumption', 'value': text} for text in report.get('assumptions', []))
    source_sheet = add_sheet(workbook, 'Source Notes', [('field', 'Field'), ('value', 'Value')], metadata,
              'Current reference snapshot, not a reconstruction of historical costs or analyst decisions.')
    add_source_note_conditional_formatting(source_sheet)
    return workbook, {'prca': prca, 'requests': len(requests), 'requested_lines': len(representatives),
                      'excluded_unpriced_lines': len(excluded_unpriced),
                      'tier_scenarios': len(scenarios), 'lines_with_impact': available}


def highlight_rows(sheet, records, current_field, candidate_field, suggested_field, analyst_tier):
    colors = {'current': 'DDEBF7', 'analyst': 'FFF2CC', 'suggested': 'E2F0D9'}
    for row_number, record in enumerate(records, start=4):
        tier = record.get(candidate_field)
        color = ('suggested' if record.get(suggested_field) else
                 'analyst' if tier == analyst_tier else
                 'current' if tier == record.get(current_field) else None)
        if color:
            for column in range(1, sheet.max_column + 1):
                sheet.cell(row_number, column).fill = PatternFill('solid', fgColor=colors[color])


def add_item_level_conditional_formatting(sheet):
    headers = {str(cell.value): cell.column_letter for cell in sheet[3] if cell.value}
    if sheet.max_row < 4 or not headers:
        return
    candidate = headers.get('Candidate Tier') or headers.get('Tier')
    if not candidate:
        return
    first_data_row, last_data_row = 4, sheet.max_row
    last_column = sheet.cell(3, sheet.max_column).column_letter
    target_range = f'A{first_data_row}:{last_column}{last_data_row}'
    exception_terms = []
    for label in ('Validation Exceptions', 'Screening Flags', 'Exceptions'):
        column = headers.get(label)
        if column:
            exception_terms.append(f'${column}{first_data_row}<>""')
    tier_ten_flag = headers.get('Tier 10 Review Flag')
    if tier_ten_flag:
        exception_terms.append(f'${tier_ten_flag}{first_data_row}=TRUE')
    price_cost = headers.get('Price > Cost') or headers.get('Price Exceeds Tier Cost')
    if price_cost:
        exception_terms.append(f'${price_cost}{first_data_row}=FALSE')
    if exception_terms:
        sheet.conditional_formatting.add(target_range, FormulaRule(
            formula=[f'AND(${candidate}{first_data_row}<>"",OR({",".join(exception_terms)}))'],
            fill=PatternFill('solid', fgColor='FCE4D6'), stopIfTrue=True))

    request_suggestion = headers.get('Request-price Suggested Tier')
    if request_suggestion:
        sheet.conditional_formatting.add(target_range, FormulaRule(
            formula=[f'AND(${candidate}{first_data_row}=${request_suggestion}{first_data_row},${request_suggestion}{first_data_row}<>"")'],
            fill=PatternFill('solid', fgColor='D9EAD3'), stopIfTrue=True))

    category_suggestion = headers.get('Category Suggested Tier') or headers.get('Suggested Tier')
    if category_suggestion:
        sheet.conditional_formatting.add(target_range, FormulaRule(
            formula=[f'AND(${candidate}{first_data_row}=${category_suggestion}{first_data_row},${category_suggestion}{first_data_row}<>"")'],
            fill=PatternFill('solid', fgColor='D9EAD3'), stopIfTrue=True))

    benchmark_met = (headers.get('Price Meets Benchmark') or
                     headers.get('Price Meets Tier Benchmark') or
                     headers.get('Price Meets Target Benchmark'))
    if benchmark_met:
        sheet.conditional_formatting.add(target_range, FormulaRule(
            formula=[f'${benchmark_met}{first_data_row}=TRUE'],
            fill=PatternFill('solid', fgColor='E2F0D9'), stopIfTrue=True))

    current_tier = headers.get('Current Tier')
    if current_tier:
        sheet.conditional_formatting.add(target_range, FormulaRule(
            formula=[f'${candidate}{first_data_row}=${current_tier}{first_data_row}'],
            fill=PatternFill('solid', fgColor='DDEBF7'), stopIfTrue=True))

    analyst_tier = (headers.get('Analyst Tier') or headers.get('Analyst Tier 7') or
                    headers.get('Analyst Benchmark Tier'))
    if analyst_tier:
        sheet.conditional_formatting.add(target_range, FormulaRule(
            formula=[f'AND(${candidate}{first_data_row}=${analyst_tier}{first_data_row},${analyst_tier}{first_data_row}<>"")'],
            fill=PatternFill('solid', fgColor='FFF2CC'), stopIfTrue=True))


def color_benchmark_rows(sheet, records):
    add_item_level_conditional_formatting(sheet)


def add_review_row_conditional_formatting(sheet):
    headers = {str(cell.value): cell.column_letter for cell in sheet[3] if cell.value}
    if sheet.max_row < 4:
        return
    target = f'A4:{sheet.cell(3, sheet.max_column).column_letter}{sheet.max_row}'
    exception_columns = [headers[label] for label in
        ('Validation Exceptions', 'Impact Exceptions', 'Exceptions') if label in headers]
    if exception_columns:
        expression = ','.join(f'${column}4<>""' for column in exception_columns)
        sheet.conditional_formatting.add(target, FormulaRule(
            formula=[f'OR({expression})'], fill=PatternFill('solid', fgColor='FCE4D6'), stopIfTrue=True))
    suggested = headers.get('Suggested for Review')
    if suggested:
        sheet.conditional_formatting.add(target, FormulaRule(
            formula=[f'${suggested}4=TRUE'], fill=PatternFill('solid', fgColor='D9EAD3'), stopIfTrue=True))
    screening = headers.get('Screening Flags')
    if screening:
        sheet.conditional_formatting.add(target, FormulaRule(
            formula=[f'${screening}4<>""'], fill=PatternFill('solid', fgColor='FFF2CC'), stopIfTrue=True))
    current_tier, candidate_tier = headers.get('Current Tier'), headers.get('Candidate Tier')
    if current_tier and candidate_tier:
        sheet.conditional_formatting.add(target, FormulaRule(
            formula=[f'${candidate_tier}4=${current_tier}4'],
            fill=PatternFill('solid', fgColor='DDEBF7'), stopIfTrue=True))
    analyst_tier = headers.get('Analyst Benchmark Tier')
    if analyst_tier and candidate_tier:
        sheet.conditional_formatting.add(target, FormulaRule(
            formula=[f'${candidate_tier}4=${analyst_tier}4'],
            fill=PatternFill('solid', fgColor='FFF2CC'), stopIfTrue=True))


def add_tier_matrix_conditional_formatting(sheet, direction: str):
    if sheet.max_row < 3 or direction not in ('cost', 'margin'):
        return
    tier_columns = [cell.column for cell in sheet[3]
                    if str(cell.value or '').strip().lower().startswith('tier ')]
    if not tier_columns:
        return
    first, last = min(tier_columns), max(tier_columns)
    last_row = max(sheet.max_row, 203)
    colors = (('D9EAD3', 'FFF2CC', 'FCE4D6') if direction == 'cost'
              else ('FCE4D6', 'FFF2CC', 'D9EAD3'))
    sheet.conditional_formatting.add(
        f'{sheet.cell(4, first).column_letter}4:{sheet.cell(last_row, last).column_letter}{last_row}',
        ColorScaleRule(start_type='min', start_color=colors[0], mid_type='percentile',
            mid_value=50, mid_color=colors[1], end_type='max', end_color=colors[2]))


def add_activity_data_bars(sheet):
    headers = {str(cell.value): cell.column_letter for cell in sheet[3] if cell.value}
    if sheet.max_row < 3:
        return
    last_row = max(sheet.max_row, 203)
    for label in ('Quantity 3 Months', 'Sales 3 Months', 'Quantity 3m', 'Sales 3m'):
        column = headers.get(label)
        if column:
            sheet.conditional_formatting.add(
                f'{column}4:{column}{last_row}',
                DataBarRule(start_type='min', end_type='max', color='5B9BD5', showValue=True))


def add_context_status_conditional_formatting(sheet):
    if sheet.max_row < 4 or sheet.max_column < 2:
        return
    target = f'A4:B{sheet.max_row}'
    sheet.conditional_formatting.add(target, FormulaRule(
        formula=['AND(OR($A4="Recommendation status",$A4="Request-price recommendation status"),$B4="CALCULATED_FOR_REVIEW")'],
        fill=PatternFill('solid', fgColor='D9EAD3'), stopIfTrue=True))
    sheet.conditional_formatting.add(target, FormulaRule(
        formula=['OR($B4="REVIEW_REQUIRED",$B4="NOT_APPROVED",AND($A4="Approval state",ISNUMBER(SEARCH("not approved",LOWER($B4)))))'],
        fill=PatternFill('solid', fgColor='FFF2CC'), stopIfTrue=True))
    sheet.conditional_formatting.add(target, FormulaRule(
        formula=['AND($A4="Exceptions",$B4<>"",$B4<>"[]")'],
        fill=PatternFill('solid', fgColor='FCE4D6'), stopIfTrue=True))


def add_analyst_note_conditional_formatting(sheet):
    headers = {str(cell.value): cell.column_letter for cell in sheet[3] if cell.value}
    decision = headers.get('Review Observation')
    if decision and sheet.max_row >= 4:
        target = f'{decision}4:{decision}{sheet.max_row}'
        for value, color in (('Approve', 'D9EAD3'), ('Reject', 'FCE4D6'),
                             ('Needs Review', 'FFF2CC')):
            sheet.conditional_formatting.add(target, FormulaRule(
                formula=[f'${decision}4="{value}"'],
                fill=PatternFill('solid', fgColor=color), stopIfTrue=True))


def add_source_note_conditional_formatting(sheet):
    if sheet.max_row >= 4:
        sheet.conditional_formatting.add(f'A4:B{sheet.max_row}', FormulaRule(
            formula=['ISNUMBER(SEARCH("assumption",$A4))'],
            fill=PatternFill('solid', fgColor='FFF2CC'), stopIfTrue=True))


def export_customers(report: dict, output_dir: Path) -> dict:
    if report.get('mode') != 'INTERNAL_REVIEW_ONLY_NOT_APPROVED_NOT_LOADABLE':
        raise ValueError('Input must be an internal review report')
    grouped = split_customers(report)
    if output_dir.exists():
        raise ValueError('Choose a new output directory; existing files are never overwritten')
    output_dir.mkdir(parents=True, exist_ok=False)
    index_rows = []
    for prca, scenarios in grouped.items():
        workbook, summary = make_customer_workbook(prca, scenarios, report)
        filename = f'Customer_{prca}_Review.xlsx'
        try:
            workbook.save(output_dir / filename)
        finally:
            workbook.close()
        index_rows.append(dict(summary, filename=filename))
    index = Workbook()
    index.remove(index.active)
    sheet = add_sheet(index, 'Customers', [
        ('prca', 'Customer PRCA'), ('requests', 'Requests'), ('requested_lines', 'Requested Lines'),
        ('excluded_unpriced_lines', 'Unpriced Lines Excluded'),
        ('tier_scenarios', 'Tier Scenarios'), ('lines_with_impact', 'Lines With Impact'),
        ('filename', 'Open Customer Workbook'),
    ], index_rows, 'Internal customer workbook index. Requested-item review only; no approvals or production loads.')
    for row_number, record in enumerate(index_rows, start=4):
        sheet.cell(row_number, 7).hyperlink = record['filename']
        sheet.cell(row_number, 7).style = 'Hyperlink'
    index.save(output_dir / 'Customer_Index.xlsx')
    index.close()
    manifest = {'exported_at': utc_now(), 'customers': len(index_rows),
                'requested_lines': sum(row['requested_lines'] for row in index_rows),
                'excluded_unpriced_lines': len(report.get('excluded_unpriced_lines', [])),
                'tier_scenarios': sum(row['tier_scenarios'] for row in index_rows),
                'source_hash': report.get('source_hash'), 'reference_hash': report.get('reference_hash'),
                'coverage': report.get('coverage', 'REQUESTED_ITEMS_ONLY'),
                'analyst_tier_assumption': report.get('analyst_tier_assumption'),
                'batch_metadata': report.get('batch_metadata'), 'files': index_rows}
    with (output_dir / 'manifest.json').open('x', encoding='utf-8') as destination:
        json.dump(manifest, destination, indent=2)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True, help='Internal review.json')
    parser.add_argument('--output-dir', type=Path, required=True)
    args = parser.parse_args()
    report = json.loads(args.input.read_text(encoding='utf-8'))
    manifest = export_customers(report, args.output_dir)
    print(json.dumps({'customers': manifest['customers'], 'requested_lines': manifest['requested_lines'],
                      'tier_scenarios': manifest['tier_scenarios'], 'output_directory': str(args.output_dir)}, indent=2))


if __name__ == '__main__':
    main()