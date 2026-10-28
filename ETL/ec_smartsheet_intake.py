"""Extract Smartsheet rebate requests into the isolated sandbox pipeline tables."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlparse
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from openpyxl import load_workbook

from ec_etl_preview import connect_from_environment, json_value
from ec_workflow import canonical_json, utc_now

API_ROOT = 'https://api.smartsheet.com/2.0'
SHEET_STAGES = (
    ('landing', '727058820845444'),
    ('intermediate', '6992940959223684'),
    ('completed', '7801368878206852'),
    ('rejected', '4985786021662596'),
)
TARGET_ZONE = ZoneInfo('America/New_York')
MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_SHEET_PAGE_SIZE = 500
HEADER_ALIASES = {
    'request_id': ('ORIGINAL ROW ID', 'ROW ID'),
    'account_number': ('CUSTOMER ACCOUNT NUMBER/PRCA', 'PRCA'),
    'created_date': ('CREATED', 'CREATED DATE'),
    'requester': ('REQUESTER NAME',),
    'notes': ('ADDITIONAL NOTES FOR AWARENESS',),
}
RAW_HEADER_TABLE = 'PRD_SBX_MMS_DB.PRC.EC_REBATE_REQUEST_HEADER_ETL_TEST'
RAW_LINE_TABLE = 'PRD_SBX_MMS_DB.PRC.EC_REBATE_REQUEST_LINES_ETL_TEST'
RAW_ATTACHMENT_TABLE = 'PRD_SBX_MMS_DB.PRC.EC_REBATE_REQUEST_ATTACHMENTS_ETL_TEST'
RAW_HEADER_FIELDS = (
    'row_id', 'created_dt', 'cmpltd_flg', 'cmpltd_dt', 'bus_pltfrm_dsc',
    'est_opp', 'cust_num_prca', 'addtl_bt_num', 'addtl_bt_num_raw', 'analyst', 'reqeust_by',
    'request_email', 'sls_rep_email', 'cntrct_type', 'cust_type', 'acct_to_mirror',
    'snd_to_cntrct_flg', 'acct_to_mirror_raw', 'source_stage', 'source_sheet_id', 'source_sheet_row_id',
    'requester_name', 'additional_notes', 'source_header_json', 'attachment_count',
    'ingestion_run_id', 'extracted_at',
)
RAW_LINE_FIELDS = (
    'row_id', 'item_e1_num', 'item_dsc', 'req_uom', 'req_price', 'req_price_raw',
    'source_attachment_id', 'source_attachment_name', 'source_attachment_sha256',
    'source_sheet_name', 'source_row_number', 'ingestion_run_id', 'extracted_at',
)
RAW_ATTACHMENT_FIELDS = (
    'row_id', 'source_stage', 'source_sheet_id', 'source_sheet_row_id',
    'attachment_id', 'attachment_name', 'attachment_type', 'file_extension',
    'attachment_size_bytes', 'content_sha256', 'download_status', 'parse_status',
    'parsed_line_count', 'error_reason', 'ingestion_run_id', 'extracted_at',
)


def parse_request_pairs(values: list[str]) -> dict[str, str]:
    targets: dict[str, str] = {}
    for value in values:
        request_id, separator, account = value.partition('=')
        request_id = request_id.strip()
        account = account.strip()
        if not separator or not request_id or not account or not account.isdigit():
            raise ValueError('Each --request must be REQUEST_ID=NUMERIC_ACCOUNT')
        if request_id in targets and targets[request_id] != account:
            raise ValueError(f'Conflicting accounts supplied for request {request_id}')
        targets[request_id] = account
    return targets


def normalize_title(value: object) -> str:
    return ' '.join(re.sub(r'[^A-Z0-9]+', ' ', str(value or '').strip().upper()).split())


def find_line_columns(values: tuple | list) -> dict[str, int] | None:
    titles = [normalize_title(value) for value in values]

    def matches(exact: set[str], predicate) -> list[int]:
        return [index for index, title in enumerate(titles)
                if title in exact or predicate(title)]

    item_columns = matches(
        {'E1 ITEM NUMBER', 'ITEM E1 NUM', 'ITEM NUMBER'},
        lambda title: 'ITEM' in title and ('E1' in title or 'NUMBER' in title),
    )
    uom_columns = matches(
        {'UOM', 'REQ UOM', 'REQUESTED UOM', 'UOM PROVIDED', 'SELL UOM'},
        lambda title: title.endswith(' UOM') or title == 'UOM',
    )
    price_columns = matches(
        {'FUTURE SELL PRICE', 'REQUESTED SELL PRICE', 'REQUESTED PRICE',
         'REQ PRICE', 'FUTURE PRICE', 'SELL PRICE'},
        lambda title: title == 'SELL PRICE',
    )
    if not item_columns or not uom_columns or not price_columns:
        return None
    if len(item_columns) != 1 or len(uom_columns) != 1 or len(price_columns) != 1:
        raise ValueError(
            'Attachment has missing or ambiguous item, UOM, or price columns: '
            + ', '.join(titles)
        )

    description_columns = [index for index, title in enumerate(titles)
                           if title in {'ITEM DESCRIPTION', 'ITEM DESC', 'DESCRIPTION'}]
    if len(description_columns) > 1:
        raise ValueError('Attachment has ambiguous item description columns')
    return {
        'item': item_columns[0],
        'uom': uom_columns[0],
        'price': price_columns[0],
        'description': description_columns[0] if description_columns else -1,
    }


def build_line_record(values: tuple | list, columns: dict[str, int], request_id: str,
                      source_file: str, source_sheet: str, source_row: int) -> dict:
    def value_for(field: str):
        index = columns[field]
        return values[index] if 0 <= index < len(values) else None

    return {
        'request_id': request_id,
        'item_e1_num': value_for('item'),
        'item_dsc': value_for('description'),
        'uom_provided': value_for('uom'),
        'future_sell_price': value_for('price'),
        'source_file': source_file,
        'source_sheet': source_sheet,
        'source_row_number': source_row,
    }


def parse_xlsx_attachment(content: bytes, request_id: str, filename: str) -> list[dict]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parsed_lines: list[dict] = []
    try:
        for worksheet in workbook.worksheets:
            header = None
            header_row = None
            for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if row_number > 100:
                    break
                columns = find_line_columns(values)
                if columns:
                    header = columns
                    header_row = row_number
                    break
            if header is None or header_row is None:
                continue
            for row_number, values in enumerate(
                    worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                    start=header_row + 1):
                indices = [header['item'], header['uom'], header['price']]
                if all(index >= len(values) or values[index] is None
                       or not str(values[index]).strip() for index in indices):
                    continue
                parsed_lines.append(build_line_record(
                    values, header, request_id, filename, worksheet.title, row_number,
                ))
    finally:
        workbook.close()
    if not parsed_lines:
        raise ValueError(f'No supported line-item rows found in {filename}')
    return parsed_lines


def parse_csv_attachment(content: bytes, request_id: str, filename: str) -> list[dict]:
    text = content.decode('utf-8-sig')
    rows = list(csv.reader(io.StringIO(text)))
    for header_index, values in enumerate(rows[:100]):
        columns = find_line_columns(values)
        if not columns:
            continue
        parsed_lines = []
        for row_index, line_values in enumerate(rows[header_index + 1:], start=header_index + 2):
            indices = [columns['item'], columns['uom'], columns['price']]
            if all(index >= len(line_values) or not str(line_values[index]).strip()
                   for index in indices):
                continue
            parsed_lines.append(build_line_record(
                line_values, columns, request_id, filename, 'CSV', row_index,
            ))
        if parsed_lines:
            return parsed_lines
    raise ValueError(f'No supported line-item rows found in {filename}')


def parse_attachment(content: bytes, request_id: str, filename: str) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    if suffix == '.xlsx':
        return parse_xlsx_attachment(content, request_id, filename)
    if suffix == '.csv':
        return parse_csv_attachment(content, request_id, filename)
    raise ValueError(f'Unsupported attachment type: {suffix or "no extension"}')


def parse_created_date(value: object) -> date:
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value or '').strip()
        if not text:
            raise ValueError('Request is missing its Created date')
        try:
            parsed = datetime.fromisoformat(text.replace('Z', '+00:00'))
        except ValueError as error:
            raise ValueError(f'Unrecognized Created date: {text}') from error
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(TARGET_ZONE)
    return parsed.date()


def extract_cell_map(row: dict, column_map: dict[int, str]) -> dict[str, object]:
    values = {}
    for cell in row.get('cells', []):
        title = column_map.get(cell.get('columnId'))
        if title:
            value = cell.get('value')
            values[title] = value if value is not None else cell.get('displayValue')
    return values


def field_value(cells: dict[str, object], aliases: tuple[str, ...]) -> object:
    normalized = {normalize_title(title): value for title, value in cells.items()}
    values = [normalized[normalize_title(alias)] for alias in aliases
              if normalize_title(alias) in normalized
              and normalized[normalize_title(alias)] is not None
              and str(normalized[normalize_title(alias)]).strip()]
    if not values:
        return None
    distinct = {str(value).strip() for value in values}
    if len(distinct) > 1:
        raise ValueError(f'Conflicting Smartsheet fields for {aliases[0]}')
    return values[0]


def get_account(cells: dict[str, object]) -> str:
    value = field_value(cells, HEADER_ALIASES['account_number'])
    if value is None:
        raise ValueError('Smartsheet row is missing Customer Account Number/PRCA')
    try:
        account = Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError('Smartsheet account must be numeric') from error
    if not account.is_finite() or account <= 0 or account != account.to_integral_value():
        raise ValueError('Smartsheet account must be a positive whole number')
    return format(account, 'f').split('.')[0]


def select_sheet_rows(sheet_rows: list[dict], stage: str, column_map: dict[int, str],
                      targets: dict[str, str], since: date | None,
                      until: date | None) -> list[dict]:
    selected = []
    for row in sheet_rows:
        cells = extract_cell_map(row, column_map)
        request_value = None
        for alias in HEADER_ALIASES['request_id']:
            request_value = field_value(cells, (alias,))
            if request_value is not None:
                break
        if request_value is None:
            continue
        request_id = str(request_value).strip()
        if targets and request_id not in targets:
            continue
        created = parse_created_date(field_value(cells, HEADER_ALIASES['created_date']))
        if since and created < since:
            continue
        if until and created >= until:
            continue
        try:
            account = get_account(cells)
        except ValueError:
            if request_id in targets:
                raise
            account = None
        if request_id in targets and targets[request_id] != account:
            raise ValueError(f'Request {request_id} account does not match supplied account')
        selected.append({
            'request_id': request_id,
            'account_number': account,
            'created_date': created.isoformat(),
            'requester_name': field_value(cells, HEADER_ALIASES['requester']),
            'additional_notes_for_awareness': field_value(cells, HEADER_ALIASES['notes']),
            'source_stage': stage,
            'source_sheet_row_id': str(row.get('id')),
            'source_header': cells,
        })
    return selected


class SmartsheetClient:
    def __init__(self, token: str) -> None:
        self.token = token

    def get_json(self, path: str, params: dict | None = None) -> dict:
        query = f'?{urlencode(params)}' if params else ''
        request = Request(
            f'{API_ROOT}/{path.lstrip("/")}{query}',
            headers={'Authorization': f'Bearer {self.token}', 'Accept': 'application/json'},
        )
        try:
            with urlopen(request, timeout=90) as response:
                return json.loads(response.read())
        except (HTTPError, URLError, TimeoutError) as error:
            raise RuntimeError('Smartsheet API request failed') from error

    def download(self, url: str, request_row_id: str) -> bytes:
        if urlparse(url).scheme.lower() != 'https':
            raise ValueError('Attachment download URL must use HTTPS')
        request = Request(url, headers={'Accept': 'application/octet-stream'})
        try:
            with urlopen(request, timeout=90) as response:
                declared_size = response.headers.get('Content-Length')
                if declared_size and int(declared_size) > MAX_ATTACHMENT_BYTES:
                    raise ValueError('Attachment exceeds the 20 MB size limit')
                content = response.read(MAX_ATTACHMENT_BYTES + 1)
        except (HTTPError, URLError, TimeoutError) as error:
            raise RuntimeError(f'Attachment download failed for row {request_row_id}') from error
        if len(content) > MAX_ATTACHMENT_BYTES:
            raise ValueError('Attachment exceeds the 20 MB size limit')
        if not content:
            raise ValueError(f'Attachment is empty for Smartsheet row {request_row_id}')
        return content


def api_get(client: SmartsheetClient, path: str, params: dict | None = None) -> dict:
    return client.get_json(path, params)


def fetch_sheet(client: SmartsheetClient, sheet_id: str) -> tuple[dict[int, str], list[dict]]:
    columns: dict[int, str] = {}
    rows: list[dict] = []
    page = 1
    while True:
        payload = api_get(client, f'sheets/{sheet_id}',
                          {'page': page, 'pageSize': MAX_SHEET_PAGE_SIZE})
        if page == 1:
            columns = {int(column['id']): str(column['title'])
                       for column in payload.get('columns', [])}
        rows.extend(payload.get('rows', []))
        total_pages = int(payload.get('totalPages') or page)
        if page >= total_pages:
            break
        page += 1
    return columns, rows


def download_attachment(client: SmartsheetClient, sheet_id: str,
                        row_id: str, attachment: dict) -> bytes:
    details = api_get(client, f'sheets/{sheet_id}/attachments/{attachment["id"]}')
    url = str(details.get('url') or '')
    return client.download(url, row_id)


def fetch_request_data(token: str, targets: dict[str, str], since: date | None,
                       until: date | None,
                       stages: tuple[tuple[str, str], ...] = SHEET_STAGES
                       ) -> tuple[list[dict], list[dict], list[dict], list[dict]]:
    client = SmartsheetClient(token)
    selected_requests: list[dict] = []
    try:
        for stage, sheet_id in stages:
            print(f'Reading Smartsheet stage: {stage}', flush=True)
            columns, sheet_rows = fetch_sheet(client, sheet_id)
            stage_requests = select_sheet_rows(
                sheet_rows, stage, columns, targets, since, until,
            )
            print(f'Found {len(stage_requests)} matching requests in {stage}', flush=True)
            selected_requests.extend(stage_requests)

        selected_by_id: dict[str, list[dict]] = defaultdict(list)
        for request in selected_requests:
            selected_by_id[request['request_id']].append(request)
        stage_rank = {stage: index for index, (stage, _) in enumerate(SHEET_STAGES)}
        duplicate_stage_rows = {}
        for request_id, requests in selected_by_id.items():
            if len(requests) == 1:
                selected_by_id[request_id] = requests[0]
                continue
            stages_for_request = [row['source_stage'] for row in requests]
            if len(set(stages_for_request)) != len(stages_for_request):
                raise ValueError(f'Duplicate request rows within a Smartsheet stage: {request_id}')
            accounts = {row['account_number'] for row in requests if row['account_number'] is not None}
            created_dates = {row['created_date'] for row in requests}
            if len(accounts) > 1 or len(created_dates) != 1:
                raise ValueError(f'Conflicting account/date across stages for request {request_id}')
            selected = max(requests, key=lambda row: stage_rank[row['source_stage']]).copy()
            if selected['account_number'] is None and accounts:
                selected['account_number'] = next(iter(accounts))
            selected['duplicate_source_stages'] = sorted(stages_for_request)
            selected_by_id[request_id] = selected
            duplicate_stage_rows[request_id] = sorted(stages_for_request)
        if targets:
            missing_ids = sorted(set(targets) - set(selected_by_id))
            if missing_ids:
                raise ValueError(f'Requests not found in selected Smartsheet stages/date range: {missing_ids}')

        request_headers = []
        request_lines = []
        request_attachments = []
        for request_id in sorted(selected_by_id):
            request = selected_by_id[request_id]
            print(f'Reading attachments for {request_id}', flush=True)
            sheet_id = next(sheet_id for stage, sheet_id in stages
                            if stage == request['source_stage'])
            attachments = api_get(
                client,
                f'sheets/{sheet_id}'
                f'/rows/{request["source_sheet_row_id"]}/attachments',
            ).get('data', [])
            attachment_hashes = set()
            request_line_count = 0
            attachment_metadata = []
            attachment_errors = []
            for attachment in attachments:
                filename = str(attachment.get('name') or 'attachment')
                extension = Path(filename).suffix.lower()
                size_kb = attachment.get('sizeInKb')
                try:
                    size_bytes = int(float(size_kb) * 1024) if size_kb is not None else None
                except (TypeError, ValueError):
                    size_bytes = None
                attachment_record = {
                    'request_id': request_id,
                    'source_stage': request['source_stage'],
                    'source_sheet_id': sheet_id,
                    'source_sheet_row_id': request['source_sheet_row_id'],
                    'attachment_id': str(attachment.get('id') or ''),
                    'attachment_name': filename,
                    'attachment_type': attachment.get('attachmentType'),
                    'file_extension': extension,
                    'attachment_size_bytes': size_bytes,
                    'content_sha256': None,
                    'download_status': 'SKIPPED',
                    'parse_status': 'NOT_ATTEMPTED',
                    'parsed_line_count': 0,
                    'error_reason': None,
                }
                if extension not in {'.xlsx', '.csv'}:
                    attachment_record['parse_status'] = 'UNSUPPORTED'
                    attachment_record['error_reason'] = 'Unsupported attachment type'
                    request_attachments.append(attachment_record)
                    continue
                if not request['account_number']:
                    attachment_record['parse_status'] = 'SKIPPED_HEADER_MISSING_PRCA'
                    attachment_record['error_reason'] = 'PRCA/account missing on request header'
                    request_attachments.append(attachment_record)
                    continue
                try:
                    content = download_attachment(
                        client, sheet_id,
                        request['source_sheet_row_id'], attachment,
                    )
                except Exception as error:
                    attachment_record['download_status'] = 'FAILED'
                    attachment_record['error_reason'] = f'Download failed: {type(error).__name__}'
                    request_attachments.append(attachment_record)
                    attachment_errors.append({'name': filename,
                                              'reason': f'Download failed: {type(error).__name__}'})
                    continue
                content_hash = hashlib.sha256(content).hexdigest()
                attachment_record['content_sha256'] = content_hash
                attachment_record['download_status'] = 'SUCCEEDED'
                if content_hash in attachment_hashes:
                    attachment_record['parse_status'] = 'DUPLICATE_CONTENT'
                    attachment_record['error_reason'] = 'Duplicate attachment content'
                    request_attachments.append(attachment_record)
                    attachment_errors.append({'name': filename, 'reason': 'Duplicate attachment content'})
                    continue
                attachment_hashes.add(content_hash)
                try:
                    lines = parse_attachment(content, request_id, filename)
                except Exception as error:
                    attachment_record['parse_status'] = 'FAILED'
                    attachment_record['error_reason'] = f'Parse failed: {type(error).__name__}: {error}'
                    request_attachments.append(attachment_record)
                    attachment_errors.append({'name': filename,
                                              'reason': f'Parse failed: {type(error).__name__}: {error}'})
                    continue
                for line in lines:
                    line.update({
                        'account_number': request['account_number'],
                        'created_date': request['created_date'],
                        'source_stage': request['source_stage'],
                        'source_sheet_id': sheet_id,
                        'source_sheet_row_id': request['source_sheet_row_id'],
                        'attachment_id': attachment_record['attachment_id'],
                        'attachment_sha256': content_hash,
                    })
                attachment_record['parse_status'] = 'PARSED'
                attachment_record['parsed_line_count'] = len(lines)
                request_attachments.append(attachment_record)
                request_lines.extend(lines)
                request_line_count += len(lines)
                attachment_metadata.append({
                    'name': filename,
                    'attachment_id': attachment_record['attachment_id'],
                    'sha256': content_hash,
                    'line_count': len(lines),
                })
            request_headers.append({
                **request,
                'source_sheet_id': sheet_id,
                'attachments': [item for item in request_attachments
                                if item['request_id'] == request_id],
                'attachment_errors': attachment_errors,
                'duplicate_source_stages': request.get('duplicate_source_stages', []),
                'line_count': request_line_count,
            })
        return request_headers, request_lines, request_attachments, selected_requests
    finally:
        client.token = ''


def source_value(cells: dict[str, object], title: str):
    return cells.get(title)


def build_raw_header(header: dict, ingestion_run_id: str, extracted_at: datetime,
                     attachment_count: int) -> dict:
    cells = header.get('source_header', {})
    additional_bts_raw = source_value(cells, 'Additional BTs')
    additional_bts_number = None
    if additional_bts_raw is not None:
        try:
            parsed_bts = Decimal(str(additional_bts_raw).strip())
            if parsed_bts.is_finite() and parsed_bts >= 0:
                additional_bts_number = parsed_bts
        except (InvalidOperation, ValueError, TypeError):
            pass
    account_to_mirror_raw = source_value(cells, 'Account to Mirror')
    account_to_mirror_number = None
    if account_to_mirror_raw is not None:
        try:
            parsed_account = Decimal(str(account_to_mirror_raw).strip())
            if parsed_account.is_finite() and parsed_account >= 0:
                account_to_mirror_number = parsed_account
        except (InvalidOperation, ValueError, TypeError):
            pass
    return {
        'row_id': header['request_id'],
        'created_dt': header['created_date'],
        'cmpltd_flg': source_value(cells, 'Completed'),
        'cmpltd_dt': source_value(cells, 'Completed Date'),
        'bus_pltfrm_dsc': source_value(cells, 'Sales Segment'),
        'est_opp': source_value(cells, 'Estimate McKesson Brands Opportunity'),
        'cust_num_prca': header.get('account_number'),
        'addtl_bt_num': additional_bts_number,
        'addtl_bt_num_raw': None if additional_bts_raw is None else str(additional_bts_raw),
        'analyst': source_value(cells, 'Assigned Analyst'),
        'reqeust_by': source_value(cells, 'Requested by'),
        'request_email': source_value(cells, 'Requester Email'),
        'sls_rep_email': source_value(cells, 'Sales Representative Email'),
        'cntrct_type': source_value(cells, 'Contract Type'),
        'cust_type': source_value(cells, 'Contract Need'),
        'acct_to_mirror': account_to_mirror_number,
        'acct_to_mirror_raw': None if account_to_mirror_raw is None else str(account_to_mirror_raw),
        'snd_to_cntrct_flg': source_value(cells, 'Send to Contracts'),
        'source_stage': header['source_stage'],
        'source_sheet_id': header['source_sheet_id'],
        'source_sheet_row_id': header['source_sheet_row_id'],
        'requester_name': header.get('requester_name'),
        'additional_notes': header.get('additional_notes_for_awareness'),
        'source_header_json': json.dumps(cells, default=json_value, ensure_ascii=False),
        'attachment_count': attachment_count,
        'ingestion_run_id': ingestion_run_id,
        'extracted_at': extracted_at,
    }


def build_raw_line(line: dict, ingestion_run_id: str, extracted_at: datetime) -> dict:
    raw_price = line.get('future_sell_price')
    raw_price_text = None if raw_price is None else str(raw_price)
    numeric_price = None
    if raw_price_text and raw_price_text.strip():
        try:
            parsed_price = Decimal(raw_price_text)
            if parsed_price.is_finite():
                candidate = float(parsed_price)
                numeric_price = candidate if math.isfinite(candidate) else None
        except (InvalidOperation, ValueError, TypeError, OverflowError):
            pass
    return {
        'row_id': line['request_id'],
        'item_e1_num': None if line.get('item_e1_num') is None else str(line['item_e1_num']),
        'item_dsc': None if line.get('item_dsc') is None else str(line['item_dsc']),
        'req_uom': None if line.get('uom_provided') is None else str(line['uom_provided']),
        'req_price': numeric_price,
        'req_price_raw': raw_price_text,
        'source_attachment_id': line.get('attachment_id'),
        'source_attachment_name': line.get('source_file'),
        'source_attachment_sha256': line.get('attachment_sha256'),
        'source_sheet_name': line.get('source_sheet'),
        'source_row_number': line.get('source_row_number'),
        'ingestion_run_id': ingestion_run_id,
        'extracted_at': extracted_at,
    }


def build_raw_attachment(attachment: dict, ingestion_run_id: str,
                         extracted_at: datetime) -> dict:
    return {
        'row_id': attachment['request_id'],
        'source_stage': attachment['source_stage'],
        'source_sheet_id': attachment['source_sheet_id'],
        'source_sheet_row_id': attachment['source_sheet_row_id'],
        'attachment_id': attachment.get('attachment_id'),
        'attachment_name': attachment.get('attachment_name'),
        'attachment_type': attachment.get('attachment_type'),
        'file_extension': attachment.get('file_extension'),
        'attachment_size_bytes': attachment.get('attachment_size_bytes'),
        'content_sha256': attachment.get('content_sha256'),
        'download_status': attachment.get('download_status'),
        'parse_status': attachment.get('parse_status'),
        'parsed_line_count': attachment.get('parsed_line_count', 0),
        'error_reason': attachment.get('error_reason'),
        'ingestion_run_id': ingestion_run_id,
        'extracted_at': extracted_at,
    }


def append_raw_snapshot(connection, header_rows: list[dict], line_rows: list[dict],
                        attachment_rows: list[dict], ingestion_run_id: str,
                        apply: bool = False) -> dict:
    batches = (
        (RAW_HEADER_TABLE, RAW_HEADER_FIELDS, header_rows),
        (RAW_LINE_TABLE, RAW_LINE_FIELDS, line_rows),
        (RAW_ATTACHMENT_TABLE, RAW_ATTACHMENT_FIELDS, attachment_rows),
    )
    cursor = connection.cursor()
    try:
        cursor.execute('BEGIN')
        existing_counts = []
        for table, _, _ in batches:
            cursor.execute(
                f'SELECT COUNT(*) FROM {table} WHERE INGESTION_RUN_ID = %s',
                (ingestion_run_id,),
            )
            existing_counts.append(cursor.fetchone()[0])
        expected_counts = [len(records) for _, _, records in batches]
        result = {
            'planned_headers': expected_counts[0],
            'planned_lines': expected_counts[1],
            'planned_attachments': expected_counts[2],
            'applied': False,
            'target_tables': [table for table, _, _ in batches],
        }
        if any(existing_counts):
            if existing_counts != expected_counts:
                raise ValueError('Existing ingestion run is incomplete or has conflicting row counts')
            connection.rollback()
            return {**result, 'already_loaded': True, 'verified_existing_run': True}
        if not apply:
            connection.rollback()
            return {**result, 'already_loaded': False}

        for table, fields, records in batches:
            if not records:
                continue
            columns = ', '.join(field.upper() for field in fields)
            placeholders = ', '.join(['%s'] * len(fields))
            cursor.executemany(
                f'INSERT INTO {table} ({columns}) VALUES ({placeholders})',
                [tuple(record[field] for field in fields) for record in records],
            )
        for (table, _, _), expected_count in zip(batches, expected_counts):
            cursor.execute(
                f'SELECT COUNT(*) FROM {table} WHERE INGESTION_RUN_ID = %s',
                (ingestion_run_id,),
            )
            if cursor.fetchone()[0] != expected_count:
                raise ValueError(f'Post-insert count mismatch for {table}')
        connection.commit()
        return {**result, 'applied': True, 'already_loaded': False}
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()


def default_output_dir() -> Path:
    local_root = Path(os.environ.get('LOCALAPPDATA') or Path.home())
    run_id = datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')
    return local_root / 'ECRebateWorkflow' / f'smartsheet_intake_{run_id}'


def write_json(path: Path, value: object) -> None:
    with path.open('x', encoding='utf-8') as destination:
        json.dump(value, destination, default=json_value, indent=2, allow_nan=False)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--env-file', type=Path,
                        help='Private dotenv file; defaults to .env beside this script')
    parser.add_argument('--request', action='append', default=[], metavar='REQUEST_ID=ACCOUNT_NUMBER',
                        help='Target one request/account pair; repeat to target multiple requests')
    parser.add_argument('--stage', action='append', choices=[stage for stage, _ in SHEET_STAGES],
                        help='Limit scanning to this Smartsheet stage; repeat to scan several')
    parser.add_argument('--snapshot', type=Path,
                        help='Import a reviewed source_snapshot.json instead of refetching Smartsheet')
    parser.add_argument('--since', help='Inclusive Created date YYYY-MM-DD')
    parser.add_argument('--until', help='Exclusive Created date YYYY-MM-DD')
    parser.add_argument('--days', type=int, default=5,
                        help='Default rolling window in days when no request/date filter is supplied')
    parser.add_argument('--output-dir', type=Path, default=None,
                        help='Private local output directory; defaults under LOCALAPPDATA')
    parser.add_argument('--apply', action='store_true',
                        help='Append validated rows to the separate ETL sandbox test tables')
    args = parser.parse_args()
    if (args.since is None) != (args.until is None):
        parser.error('Provide both --since and --until, or neither')
    if args.days < 1 or args.days > 90:
        parser.error('--days must be between 1 and 90')
    try:
        targets = parse_request_pairs(args.request)
        since = date.fromisoformat(args.since) if args.since else None
        until = date.fromisoformat(args.until) if args.until else None
    except ValueError as error:
        parser.error(str(error))
    if since and until and since >= until:
        parser.error('--since must precede --until')
    if args.snapshot and (args.request or args.stage or args.since or args.until):
        parser.error('--snapshot cannot be combined with request, stage, or date filters')
    snapshot_manifest = {}
    if args.snapshot:
        manifest_path = args.snapshot.with_name('manifest.json')
        if manifest_path.is_file():
            snapshot_manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        since = date.fromisoformat(snapshot_manifest['since_inclusive']) \
            if snapshot_manifest.get('since_inclusive') else None
        until = date.fromisoformat(snapshot_manifest['until_exclusive']) \
            if snapshot_manifest.get('until_exclusive') else None
    if not args.snapshot and not targets and since is None:
        today = datetime.now(TARGET_ZONE).date()
        since, until = today - timedelta(days=args.days), today + timedelta(days=1)
    selected_stages = tuple(
        (stage, sheet_id) for stage, sheet_id in SHEET_STAGES
        if (snapshot_manifest.get('stages_scanned') and
            stage in set(snapshot_manifest['stages_scanned']))
        or (not args.snapshot and (not args.stage or stage in set(args.stage)))
    )

    env_file = args.env_file or Path(__file__).with_name('.env')
    if not env_file.is_file():
        parser.error(f'Environment file not found: {env_file}')
    load_dotenv(env_file, override=False)
    token = os.environ.get('SMARTSHEET_API_TOKEN', '').strip()
    if not args.snapshot and not token:
        parser.error('SMARTSHEET_API_TOKEN is missing from the environment/.env file')

    output_dir = args.output_dir or default_output_dir()
    if output_dir.exists():
        parser.error('Choose a new output directory; existing output is never overwritten')

    try:
        if args.snapshot:
            print(f'Loading reviewed source snapshot: {args.snapshot}', flush=True)
            source_payload = json.loads(args.snapshot.read_text(encoding='utf-8'))
            if source_payload.get('source') != 'Smartsheet API':
                raise ValueError('Snapshot source must be Smartsheet API')
            headers = source_payload['request_headers']
            lines = source_payload['request_lines']
            attachments = source_payload['request_attachments']
            selected_rows = headers
        else:
            print('Starting read-only Smartsheet extraction', flush=True)
            headers, lines, attachments, selected_rows = fetch_request_data(
                token, targets, since, until, selected_stages,
            )
            source_payload = {
                'source': 'Smartsheet API',
                'extracted_at': utc_now(),
                'request_headers': headers,
                'request_lines': lines,
                'request_attachments': attachments,
            }
        source_content = {
            'request_headers': headers,
            'request_lines': lines,
            'request_attachments': attachments,
        }
        source_hash = hashlib.sha256(canonical_json(source_content).encode('utf-8')).hexdigest()
        expected_source_hash = snapshot_manifest.get('source_hash')
        if expected_source_hash and expected_source_hash != source_hash:
            raise ValueError('Reviewed source snapshot does not match its manifest hash')
        ingestion_run_id = source_hash
        extracted_at = datetime.now(timezone.utc).replace(tzinfo=None)
        attachments_per_request = Counter(row['request_id'] for row in attachments)
        header_rows = [build_raw_header(header, ingestion_run_id, extracted_at,
                        attachments_per_request[header['request_id']])
                   for header in headers]
        line_rows = [build_raw_line(line, ingestion_run_id, extracted_at) for line in lines]
        attachment_rows = [build_raw_attachment(attachment, ingestion_run_id, extracted_at)
                           for attachment in attachments]
        output_dir.mkdir(parents=True, exist_ok=False)
        source_payload['ingestion_run_id'] = ingestion_run_id
        write_json(output_dir / 'source_snapshot.json', source_payload)

        connection = connect_from_environment()
        try:
            result = append_raw_snapshot(
                connection, header_rows, line_rows, attachment_rows,
                ingestion_run_id, apply=args.apply,
            )
        finally:
            connection.close()

        manifest = {
            'mode': 'SMARTSHEET_SANDBOX_INTAKE_APPLIED' if result['applied']
                    else 'SMARTSHEET_SANDBOX_INTAKE_DRY_RUN',
            'source_hash': source_hash,
            'since_inclusive': since.isoformat() if since else None,
            'until_exclusive': until.isoformat() if until else None,
            'requested_pairs': targets,
            'stages_scanned': [stage for stage, _ in selected_stages],
            'smartsheet_rows_selected': len(selected_rows),
            'request_headers_extracted': len(headers),
            'request_lines_extracted': len(lines),
            'request_attachments_extracted': len(attachments),
            'headers_missing_prca': sum(not header.get('account_number') for header in headers),
            'headers_without_attachments': sum(not header.get('attachments') for header in headers),
            'requests_with_attachment_exceptions': sum(
                bool(header.get('attachment_errors')) for header in headers
            ),
            'attachment_exceptions': [
                {'request_id': header['request_id'], **error}
                for header in headers for error in header.get('attachment_errors', [])
            ],
            'requests_in_multiple_stages': [
                {'request_id': header['request_id'],
                 'selected_stage': header['source_stage'],
                 'source_stages': header['duplicate_source_stages']}
                for header in headers if header.get('duplicate_source_stages')
            ],
            'target_headers': len(header_rows),
            'target_lines': len(line_rows),
            'target_attachments': len(attachment_rows),
            'target_tables': [
                'PRD_SBX_MMS_DB.PRC.EC_REBATE_REQUEST_HEADER_ETL_TEST',
                'PRD_SBX_MMS_DB.PRC.EC_REBATE_REQUEST_LINES_ETL_TEST',
                'PRD_SBX_MMS_DB.PRC.EC_REBATE_REQUEST_ATTACHMENTS_ETL_TEST',
            ],
            **result,
            'source_snapshot': 'source_snapshot.json',
            'email_sent': False,
            'pricing_or_approval_writes': False,
        }
        write_json(output_dir / 'manifest.json', manifest)
        print(json.dumps(manifest, indent=2))
    except Exception as error:
        if output_dir.exists():
            failure = {'mode': 'SMARTSHEET_SANDBOX_INTAKE_FAILED',
                       'error_type': type(error).__name__, 'error': str(error),
                       'source_snapshot_exists': (output_dir / 'source_snapshot.json').exists(),
                       'database_write_requested': args.apply}
            try:
                with (output_dir / 'failure.json').open('x', encoding='utf-8') as destination:
                    json.dump(failure, destination, indent=2)
            except OSError:
                pass
        raise


if __name__ == '__main__':
    main()
